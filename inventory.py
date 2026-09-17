"""Monthly AVCO inventory. Python 3.10+; openpyxl. No Excel formulas or macros.

Input is read-only unless --prepare-input is used (with a backup).
Values use Decimal; reports contain static values only.
Copyright 2026. You may use, modify and redistribute this script.
"""
from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation, getcontext
import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
import sys
import tempfile
import time

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from input_layout import SHEETS, FIELDS, OPTIONAL_HEADERS, SETTINGS, TYPES, canonical_columns
from input_layout import AZ_FIELDS, AZ_SHEETS, resolve_sheet, schema_metadata, canonical_type, canonical_setting

VERSION = "1.2.1"
getcontext().prec = 34
ZERO = Decimal(0)
EPS = Decimal("1e-18")
HEADERS = {
    "Settings": ["Setting", "Value"],
    "Products": ["SKU", "Product", "Unit"],
    "Purchases": ["Line ID", "Date", "Shipment ID", "Document", "SKU", "Quantity", "Currency", "Amount FCY", "FX to AZN"],
    "Expenses": ["Expense ID", "Document Date", "Shipment ID", "Expense Type", "Currency", "Amount FCY", "FX to AZN"],
    "Sales": ["Line ID", "Date", "Document", "Customer", "SKU", "Type", "Quantity", "Revenue AZN"],
    "Returns": ["Return ID", "Date", "Type", "Original Line ID", "Quantity", "Revenue Reversal AZN"],
}


class InvalidInput(Exception):
    def __init__(self, issues):
        self.issues = issues
        super().__init__(f"Input has {len(issues)} error(s).")


def issue(sheet, row, message, ident="", severity="ERROR"):
    return {"Severity": severity, "Sheet": sheet, "Row": row or "", "ID": ident, "Message": message}


def key(value):
    return str(value or "").strip().casefold()


def text(value, name, required=True):
    name = AZ_FIELDS.get(name, name).rstrip(' *')
    if value is None or not str(value).strip():
        if required:
            raise ValueError(f"{name} sahəsini doldurun.")
        return ""
    if isinstance(value, (datetime, date, bool)):
        raise ValueError(f"{name}: tarix və ya məntiqi dəyər deyil, mətn kodu tələb olunur.")
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def number(value, name, positive=False, allow_negative=False, optional=False):
    name = AZ_FIELDS.get(name, name).rstrip(' *')
    if value is None or value == "":
        if optional:
            return ZERO
        raise ValueError(f"{name} sahəsini doldurun; məbləğ sıfırdırsa, 0 yazın.")
    if isinstance(value, bool):
        raise ValueError(f"{name}: ədəd tələb olunur.")
    try:
        # Decimal comma is accepted as text; thousands separators are not guessed.
        raw = str(value).strip().replace("\u00a0", "").replace(" ", "")
        if "," in raw and "." not in raw:
            raw = raw.replace(",", ".")
        result = Decimal(raw)
    except (InvalidOperation, ValueError):
        raise ValueError(f"{name}: yanlış ədəd {value!r}.") from None
    if not result.is_finite():
        raise ValueError(f"{name}: sonlu ədəd tələb olunur.")
    if positive and result <= 0:
        raise ValueError(f"{name} sıfırdan böyük olmalıdır.")
    if not allow_negative and result < 0:
        raise ValueError(f"{name} mənfi ola bilməz.")
    return result


def as_date(value, name):
    name = AZ_FIELDS.get(name, name).rstrip(' *')
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%d.%m.%Y"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    raise ValueError(f"{name}: Excel tarixi, YYYY-MM-DD və ya DD.MM.YYYY daxil edin.")


def month(d):
    return d.replace(day=1)


def months(start, end):
    while start <= end:
        yield start
        start = date(start.year + (start.month == 12), 1 if start.month == 12 else start.month + 1, 1)


def read_input(path):
    """Read an input workbook from disk."""
    return read_input_bytes(Path(path).read_bytes())


def read_input_bytes(raw_bytes):
    """Read all nonempty rows from XLSX bytes, including rows outside a Table."""
    raw_bytes = bytes(raw_bytes)
    wb = load_workbook(BytesIO(raw_bytes), read_only=True, data_only=False)
    result, errors = {}, []
    try:
        meta = schema_metadata(wb)
        for name, required in {**HEADERS, 'Counterparties': ['Name']}.items():
            sheet_name = resolve_sheet(wb, name)
            if sheet_name is None:
                if name == 'Counterparties' and not meta and not any(
                        SHEETS[n] in wb.sheetnames or AZ_SHEETS[n] in wb.sheetnames for n in HEADERS):
                    continue
                errors.append(issue(name, 6, "Vacib vərəq yoxdur."))
                continue
            ws = wb[sheet_name]
            columns = canonical_columns([c.value for c in next(ws.iter_rows(min_row=6, max_row=6))])
            duplicates = {h for h in columns if h and columns.count(h) > 1}
            if duplicates:
                errors.append(issue(name, 6, "Təkrarlanan başlıqlar: " + ", ".join(sorted(duplicates))))
            missing = set(required) - set(columns)
            if missing:
                errors.append(issue(name, 6, "Çatışmayan sütunlar: " + ", ".join(sorted(missing))))
                continue
            indexes = {h: columns.index(h) for h in required + OPTIONAL_HEADERS.get(name, []) if h in columns}
            records = []
            for cells in ws.iter_rows(min_row=7):
                values = {h: cells[i].value if i < len(cells) else None for h, i in indexes.items()}
                if all(v is None or v == "" for v in values.values()):
                    continue
                row_no = cells[0].row
                for h, i in indexes.items():
                    if i < len(cells) and cells[i].data_type in ("f", "e"):
                        errors.append(issue(name, row_no, f"{h}: formulanı / xətanı hazır dəyərlə əvəz edin."))
                values["_row"] = row_no
                values['_modern'] = (meta['validation_profile'] == 'modern' if meta
                                     else sheet_name != name)
                if name == 'Settings':
                    values['Setting'] = canonical_setting(values['Setting'])
                for h in ('Type', 'Purchase Type'):
                    if h in values:
                        values[h] = canonical_type(values[h])
                records.append(values)
            result[name] = records
        if meta and not meta['validate_parties']:
            result.pop('Counterparties', None)
    finally:
        wb.close()
    if errors:
        raise InvalidInput(errors)
    return result, hashlib.sha256(raw_bytes).hexdigest()


def parse_input(raw):
    errors = []
    counterparties = {}
    for r in raw.get('Counterparties', []):
        try:
            value = text(r.get('Name'), 'Name')
            if key(value) in counterparties:
                raise ValueError('Kontragent siyahıda təkrarlanır.')
            counterparties[key(value)] = value
        except ValueError as e:
            errors.append(issue('Counterparties', r.get('_row', 0), str(e)))
    settings = {}
    for r in raw.get("Settings", []):
        k = key(r.get("Setting"))
        if k in settings:
            errors.append(issue("Settings", r["_row"], "Parametr təkrarlanır."))
        settings[k] = r.get("Value")
    try:
        start = as_date(settings.get("start month"), "Start Month")
        end = as_date(settings.get("report through"), "Report Through")
        if start.day != 1 or end.day != 1:
            raise ValueError("Başlanğıc və son tarix ayın ilk günü olmalıdır.")
        if end < start:
            raise ValueError("Hesabatın son ayı başlanğıc ayından əvvəldir.")
        company = text(settings.get("company"), "Company", False) or "Inventory"
    except ValueError as e:
        errors.append(issue("Settings", 0, str(e)))
        raise InvalidInput(errors)
    parsed = {name: [] for name in ("Products", "Purchases", "Expenses", "Sales", "Returns")}
    products, purchases, sales = {}, {}, {}
    id_maps = {name: set() for name in parsed}
    for name in parsed:
        for r in raw.get(name, []):
            row = r.get("_row", 0)
            try:
                id_col = {"Products": "SKU", "Expenses": "Expense ID", "Returns": "Return ID"}.get(name, "Line ID")
                ident = text(r.get(id_col), id_col)
                ik = key(ident)
                if ik in id_maps[name]:
                    raise ValueError(f"Kod təkrarlanır ({id_col}): {ident}. Mal siyahısında SKU, jurnalda sətir kodu unikal olmalıdır.")
                id_maps[name].add(ik)
                v = {"id": ident, "key": ik, "row": row}
                if name == "Products":
                    v.update(sku=ident, product=text(r.get("Product"), "Product"), unit=text(r.get("Unit"), "Unit"))
                    products[ik] = v
                else:
                    dc = "Document Date" if name == "Expenses" else "Date"
                    v["date"] = as_date(r.get(dc), dc)
                    v["month"] = month(v["date"])
                    if name != "Expenses" and v["month"] < start:
                        raise ValueError("Tarix uçotun başlanğıcından əvvəldir. Başlanğıc qalıq sıfırdır; bütün əməliyyat tarixçəsini daxil edin.")
                if name in ("Purchases", "Sales"):
                    sk = key(text(r.get("SKU"), "SKU"))
                    if sk not in products:
                        raise ValueError(f"SKU {r.get('SKU')} mal siyahısında yoxdur.")
                    v.update(sku=products[sk]["sku"], sk=sk, qty=number(r.get("Quantity"), "Quantity", positive=True), document=text(r.get("Document"), "Document"))
                if name in ("Purchases", "Expenses"):
                    v.update(shipment=text(r.get("Shipment ID"), "Shipment ID"), currency=text(r.get("Currency"), "Currency").upper(), amount=number(r.get("Amount FCY"), "Amount FCY", positive=name == "Purchases", allow_negative=name == "Expenses"), fx=number(r.get("FX to AZN"), "FX to AZN", positive=True))
                    v["sh"] = key(v["shipment"])
                    if v["currency"] == "AZN" and v["fx"] != 1:
                        raise ValueError("AZN üçün məzənnə 1 olmalıdır.")
                    v["base"] = v["amount"] * v["fx"]
                if name == "Purchases":
                    v['supplier'] = text(r.get('Supplier'), 'Supplier', r.get('_modern', False))
                    v['purchase_type'] = text(r.get('Purchase Type'), 'Purchase Type', r.get('_modern', False))
                    if v['purchase_type'] and v['purchase_type'] not in ('Import', 'Local'):
                        raise ValueError('Alış növü: İdxal və ya Yerli seçin.')
                    purchases[ik] = v
                if name == "Expenses":
                    v["type"] = text(r.get("Expense Type"), "Expense Type")
                    v['counterparty'] = text(r.get('Counterparty'), 'Counterparty', r.get('_modern', False))
                    v['document'] = text(r.get('Document'), 'Document', r.get('_modern', False))
                if name == "Sales":
                    typ = key(r.get("Type"))
                    if typ not in ("sale", "write-off"):
                        raise ValueError("Əməliyyat: Satış və ya Silinmə seçin.")
                    v["type"] = "Sale" if typ == "sale" else "Write-off"
                    v["customer"] = text(r.get("Customer"), "Customer", typ == "sale")
                    v["revenue"] = number(r.get("Revenue AZN"), "Revenue AZN", optional=typ == "write-off")
                    if typ == "write-off" and v["revenue"] != 0:
                        raise ValueError("Silinmə üçün satış məbləği 0 və ya boş olmalıdır.")
                    sales[ik] = v
                party = v.get('supplier') or v.get('counterparty') or v.get('customer')
                if party and 'Counterparties' in raw:
                    if key(party) not in counterparties:
                        raise ValueError(f'Kontragent {party!r} kontragentlər siyahısında yoxdur.')
                    field = 'supplier' if name == 'Purchases' else 'counterparty' if name == 'Expenses' else 'customer'
                    v[field] = counterparties[key(party)]
                if name == "Returns":
                    typ = key(r.get("Type"))
                    if typ not in ("customer", "supplier"):
                        raise ValueError("Qaytarma növü: Alıcı və ya Təchizatçı seçin.")
                    v["type"] = "Customer" if typ == "customer" else "Supplier"
                    v["original_id"] = text(r.get("Original Line ID"), "Original Line ID")
                    source = (sales if typ == "customer" else purchases).get(key(v["original_id"]))
                    if source is None:
                        raise ValueError("İlkin sətirin kodu satış və ya alış jurnalında tapılmadı.")
                    if typ == "customer" and source["type"] != "Sale":
                        raise ValueError("Alıcı qaytarması silinməyə deyil, satışa istinad etməlidir.")
                    if v["date"] < source["date"]:
                        raise ValueError("Qaytarma tarixi ilkin sənədin tarixindən əvvəldir.")
                    v.update(source=source, sk=source["sk"], sku=source["sku"], qty=number(r.get("Quantity"), "Quantity", positive=True))
                    v["revenue"] = number(r.get("Revenue Reversal AZN"), "Revenue Reversal AZN", optional=typ == "supplier")
                    if typ == "supplier" and v["revenue"] != 0:
                        raise ValueError("Təchizatçıya qaytarmada satış məbləğinin azalması 0 və ya boş olmalıdır.")
                parsed[name].append(v)
            except (ValueError, InvalidOperation) as e:
                errors.append(issue(name, row, str(e), str(r.get("SKU") or r.get("Line ID") or r.get("Return ID") or r.get("Expense ID") or "")))
    shipments = {p["sh"] for p in parsed["Purchases"]}
    shipment_types = {}
    for p in parsed['Purchases']:
        previous = shipment_types.setdefault(p['sh'], p['purchase_type'])
        if previous != p['purchase_type']:
            errors.append(issue('Purchases', p['row'], 'Eyni tədarük nömrəsində fərqli alış növləri var. Bəyannamə / e-qaimə nömrəsini yoxlayın.', p['id']))
    for e in parsed["Expenses"]:
        if e["sh"] not in shipments:
            errors.append(issue("Expenses", e["row"], "Bəyannamə / e-qaimə nömrəsi alışlarda tapılmadı.", e["id"]))
    returned_qty, returned_revenue = defaultdict(Decimal), defaultdict(Decimal)
    for r in sorted(parsed["Returns"], key=lambda r: (r["date"], r["key"])):
        k = (r["type"], r["source"]["key"])
        returned_qty[k] += r["qty"]
        returned_revenue[k] += r["revenue"]
        if returned_qty[k] > r["source"]["qty"]:
            errors.append(issue("Returns", r["row"], "Qaytarılan ümumi miqdar ilkin sətrin miqdarını aşır.", r["id"]))
        if r["type"] == "Customer" and returned_revenue[k] > r["source"]["revenue"] + Decimal("0.005"):
            errors.append(issue("Returns", r["row"], "Satış məbləğinin ümumi azalması ilkin satış məbləğini aşır.", r["id"]))
    if errors:
        raise InvalidInput(errors)
    return parsed, {"start": start, "end": end, "company": company}


def calculate(raw):
    data, config = parse_input(raw)
    errors = []
    shipment_lines, costs = defaultdict(list), defaultdict(Decimal)
    for p in data["Purchases"]:
        shipment_lines[p["sh"]].append(p)
    for e in data["Expenses"]:
        costs[e["sh"]] += e["base"]
    for sh, lines in shipment_lines.items():
        lines.sort(key=lambda p: p["key"])
        base = sum((p["base"] for p in lines), ZERO)
        allocated = ZERO
        for i, p in enumerate(lines):
            p["allocated"] = costs[sh] - allocated if i == len(lines) - 1 else costs[sh] * p["base"] / base
            allocated += p["allocated"]
            p["landed"] = p["base"] + p["allocated"]
            p["unit_cost"] = p["landed"] / p["qty"]
            if p["landed"] < 0:
                errors.append(issue("Purchases", p["row"], "Tam maya dəyəri mənfidir: xərcləri və düzəlişləri yoxlayın.", p["id"]))
    if errors:
        raise InvalidInput(errors)
    grouped = {name: defaultdict(list) for name in ("Purchases", "Sales", "Returns")}
    first_month = {}
    excluded = 0
    for name, groups in grouped.items():
        for row in data[name]:
            if row["month"] > config["end"]:
                excluded += 1
                continue
            k = (row["sk"], row["month"])
            groups[k].append(row)
            first_month[row["sk"]] = min(first_month.get(row["sk"], row["month"]), row["month"])
    monthly, details, writeoffs = [], [], []
    rates = {}
    for product in sorted(data["Products"], key=lambda p: p["key"]):
        sk = product["key"]
        if sk not in first_month:
            continue
        opening_qty = opening_value = ZERO
        for m in months(first_month[sk], config["end"]):
            ps, sales, returns = (grouped[n].get((sk, m), []) for n in ("Purchases", "Sales", "Returns"))
            purchase_qty = sum((p["qty"] for p in ps), ZERO)
            purchase_value = sum((p["landed"] for p in ps), ZERO)
            supplier = [r for r in returns if r["type"] == "Supplier"]
            customer = [r for r in returns if r["type"] == "Customer"]
            supplier_qty = sum((r["qty"] for r in supplier), ZERO)
            supplier_value = sum((r["qty"] * r["source"]["unit_cost"] for r in supplier), ZERO)
            prior = [r for r in customer if r["source"]["month"] < m]
            prior_qty = sum((r["qty"] for r in prior), ZERO)
            prior_value = sum((r["qty"] * rates[sk, r["source"]["month"]] for r in prior), ZERO)
            available_qty = opening_qty + purchase_qty + prior_qty - supplier_qty
            available_value = opening_value + purchase_value + prior_value - supplier_value
            if available_qty < 0 or available_value < -EPS or (available_qty == 0 and abs(available_value) > EPS):
                errors.append(issue("Monthly Inventory", 0, f"{m:%Y-%m}: istifadə edilə bilən qalıq yanlışdır: {available_qty} / {available_value} AZN. Alışları və qaytarmaları yoxlayın.", product["sku"]))
                break
            avco = available_value / available_qty if available_qty else ZERO
            rates[sk, m] = avco
            sale_qty = sum((s["qty"] for s in sales if s["type"] == "Sale"), ZERO)
            write_qty = sum((s["qty"] for s in sales if s["type"] == "Write-off"), ZERO)
            customer_qty = sum((r["qty"] for r in customer), ZERO)
            customer_value = sum((r["qty"] * rates[sk, r["source"]["month"]] for r in customer), ZERO)
            gross_cogs, write_cost = sale_qty * avco, write_qty * avco
            close_qty = opening_qty + purchase_qty + customer_qty - supplier_qty - sale_qty - write_qty
            close_value = opening_value + purchase_value + customer_value - supplier_value - gross_cogs - write_cost
            if close_qty < 0 or close_value < -EPS or (close_qty == 0 and abs(close_value) > EPS):
                errors.append(issue("Monthly Inventory", 0, f"{m:%Y-%m}: son qalıq mənfidir və ya uyğun gəlmir: {close_qty} / {close_value} AZN.", product["sku"]))
                break
            if abs(close_value) < EPS:
                close_value = ZERO
            revenue = sum((s["revenue"] for s in sales), ZERO) - sum((r["revenue"] for r in customer), ZERO)
            net_cogs = gross_cogs - customer_value
            monthly.append({"SKU": product["sku"], "Month": m, "Opening Qty": opening_qty, "Opening Value AZN": opening_value,
                "Purchases Qty": purchase_qty, "Purchases Value AZN": purchase_value, "Sales Qty": sale_qty, "Monthly AVCO AZN": avco,
                "Customer Return Qty": customer_qty, "Customer Return Value AZN": customer_value,
                "Supplier Return Qty": supplier_qty, "Supplier Return Value AZN": supplier_value,
                "Write-off Qty": write_qty, "Write-off Cost AZN": write_cost, "Gross Sales COGS AZN": gross_cogs,
                "Net COGS AZN": net_cogs, "Closing Qty": close_qty, "Closing Value AZN": close_value,
                "Net Revenue AZN": revenue, "Gross Profit AZN": revenue - net_cogs, "Product": product["product"], "Unit": product["unit"]})
            for s in sales:
                cost = s["qty"] * avco
                if s["type"] == "Write-off":
                    writeoffs.append({"Date": s["date"], "Line ID": s["id"], "Document": s["document"], "SKU": s["sku"], "Product": product["product"], "Quantity": s["qty"], "Unit Cost AZN": avco, "Write-off Cost AZN": cost})
                    continue
                details.append(sale_detail(s["date"], "Sale", s["id"], "", s, product, s["qty"], s["revenue"], cost, avco))
            for r in customer:
                rate = rates[sk, r["source"]["month"]]
                details.append(sale_detail(r["date"], "Customer Return", r["id"], r["source"]["id"], r["source"], product, -r["qty"], -r["revenue"], -r["qty"] * rate, rate))
            opening_qty, opening_value = close_qty, close_value
    if errors:
        raise InvalidInput(errors)
    details.sort(key=lambda r: (r["Date"], r["Line ID"].casefold()))
    summary = aggregate(details, ["Month", "SKU", "Product", "Customer"], ["Net Quantity", "Net Revenue AZN", "Net COGS AZN", "Gross Profit AZN"])
    for r in summary:
        r["Margin %"] = r["Gross Profit AZN"] / r["Net Revenue AZN"] if r["Net Revenue AZN"] > 0 else None
    overview = aggregate(monthly, ["Month"], ["Purchases Value AZN", "Net Revenue AZN", "Net COGS AZN", "Gross Profit AZN", "Write-off Cost AZN", "Closing Value AZN"])
    for r in overview:
        r["Margin %"] = r["Gross Profit AZN"] / r["Net Revenue AZN"] if r["Net Revenue AZN"] > 0 else None
        r["Gross Profit Less Write-offs AZN"] = r["Gross Profit AZN"] - r["Write-off Cost AZN"]
    receipts = [{"Date": p["date"], "Line ID": p["id"], "Shipment ID": p["shipment"], "Document": p["document"], "SKU": p["sku"], "Quantity": p["qty"], "Currency": p["currency"], "Amount FCY": p["amount"], "FX to AZN": p["fx"], "Base AZN": p["base"], "Allocated Costs AZN": p["allocated"], "Landed Value AZN": p["landed"], "Unit Cost AZN": p["unit_cost"]} for p in sorted(data["Purchases"], key=lambda p: (p["date"], p["key"])) if p["month"] <= config["end"]]
    checks = [issue("All", 0, "Giriş məlumatları və aylıq qalıqlar yoxlamadan keçdi.", severity="OK"),
              issue("Expenses", 0, f"Ümumi xərclər {sum(costs.values(), ZERO)} AZN; bölüşdürülən {sum((p['allocated'] for p in data['Purchases']), ZERO)} AZN. Bütün məlum xərclər sənəd tarixindən asılı olmayaraq daxil edilib.", severity="INFO")]
    if excluded:
        checks.append(issue("Settings", 0, f"Hesabat dövründən sonrakı əməliyyat sətrləri: {excluded}. Hesabata daxil edilməyib.", severity="INFO"))
    checks.extend(duplicate_warnings(data))
    purchase_by_id = {p['id']: p for p in data['Purchases']}
    for row in receipts:
        p = purchase_by_id[row['Line ID']]
        row.update({'Supplier': p['supplier'], 'Purchase Type': p['purchase_type']})
    shipment_summary, expense_detail = shipment_reports(data, shipment_lines, config['end'])
    return {"config": config, "tables": {"Monthly Summary": overview, "Monthly Inventory": monthly, "Sales Detail": details, "Sales Summary": summary, "Shipment Cost": receipts, 'Shipment Summary': shipment_summary, 'Shipment Expenses': expense_detail, "Write-offs": writeoffs, "Checks": checks}, "source_counts": {n: len(rows) for n, rows in data.items()}}


def duplicate_warnings(data):
    """Repeated business content is suspicious, but may be legitimate."""
    fields = {
        'Purchases': ('date', 'sh', 'document', 'supplier', 'sk', 'qty', 'currency', 'amount', 'fx'),
        'Expenses': ('date', 'sh', 'document', 'counterparty', 'type', 'currency', 'amount', 'fx'),
        'Sales': ('date', 'document', 'customer', 'sk', 'type', 'qty', 'revenue'),
        'Returns': ('date', 'type', 'original_id', 'qty', 'revenue'),
    }
    warnings = []
    for name, columns in fields.items():
        seen = {}
        for row in sorted(data[name], key=lambda r: r['key']):
            signature = tuple(key(row[c]) if isinstance(row[c], str) else row[c] for c in columns)
            if signature in seen:
                warnings.append(issue(name, row['row'], f"Mümkün təkrar əməliyyat: {seen[signature]['id']} sətri ilə eynidir. Sənədi yoxlayın; hər iki sətir hesablamaya daxil edilib.", row['id'], 'WARNING'))
            else:
                seen[signature] = row
    return warnings


def shipment_reports(data, shipment_lines, end):
    summaries, expenses = [], []
    expense_groups = defaultdict(list)
    for expense in data['Expenses']:
        expense_groups[expense['sh']].append(expense)
    for sh, lines in sorted(shipment_lines.items()):
        included = [p for p in lines if p['month'] <= end]
        if not included:
            continue
        label = lines[0]['shipment']
        summaries.append({'Shipment ID': label, 'Purchase Type': lines[0]['purchase_type'],
            'Supplier': ', '.join(sorted({p['supplier'] for p in included if p['supplier']})),
            'Base AZN': sum((p['base'] for p in included), ZERO),
            'Allocated Costs AZN': sum((p['allocated'] for p in included), ZERO),
            'Landed Value AZN': sum((p['landed'] for p in included), ZERO)})
        for e in sorted(expense_groups[sh], key=lambda e: (e['date'], e['key'])):
            expenses.append({'Shipment ID': label, 'Expense ID': e['id'], 'Document Date': e['date'],
                'Expense Type': e['type'], 'Counterparty': e['counterparty'], 'Document': e['document'],
                'Currency': e['currency'], 'Amount FCY': e['amount'], 'FX to AZN': e['fx'], 'Expense AZN': e['base']})
    return summaries, expenses


def sale_detail(d, typ, ident, original, s, product, qty, revenue, cost, rate):
    return {"Date": d, "Month": month(d), "Type": typ, "Line ID": ident, "Original Line ID": original,
            "Document": s["document"], "Customer": s["customer"], "SKU": product["sku"], "Product": product["product"],
            "Net Quantity": qty, "Net Revenue AZN": revenue, "Net COGS AZN": cost, "Unit Cost AZN": rate,
            "Gross Profit AZN": revenue - cost, "Margin %": (revenue - cost) / revenue if revenue > 0 else None}


def aggregate(rows, keys, values):
    buckets = {}
    for row in rows:
        k = tuple(row[x] for x in keys)
        if k not in buckets:
            buckets[k] = {**dict(zip(keys, k)), **{v: ZERO for v in values}}
        for v in values:
            buckets[k][v] += row[v]
    return [buckets[k] for k in sorted(buckets)]


EMPTY_HEADERS = {
    "Monthly Summary": ["Month", "Purchases Value AZN", "Net Revenue AZN", "Net COGS AZN", "Gross Profit AZN", "Write-off Cost AZN", "Closing Value AZN", "Margin %", "Gross Profit Less Write-offs AZN"],
    "Monthly Inventory": ["SKU", "Month", "Opening Qty", "Opening Value AZN", "Purchases Qty", "Purchases Value AZN", "Sales Qty", "Monthly AVCO AZN", "Customer Return Qty", "Customer Return Value AZN", "Supplier Return Qty", "Supplier Return Value AZN", "Write-off Qty", "Write-off Cost AZN", "Gross Sales COGS AZN", "Net COGS AZN", "Closing Qty", "Closing Value AZN", "Net Revenue AZN", "Gross Profit AZN", "Product", "Unit"],
    "Sales Detail": ["Date", "Month", "Type", "Line ID", "Original Line ID", "Document", "Customer", "SKU", "Product", "Net Quantity", "Net Revenue AZN", "Net COGS AZN", "Unit Cost AZN", "Gross Profit AZN", "Margin %"],
    "Sales Summary": ["Month", "SKU", "Product", "Customer", "Net Quantity", "Net Revenue AZN", "Net COGS AZN", "Gross Profit AZN", "Margin %"],
    "Shipment Cost": ["Date", "Line ID", "Shipment ID", "Document", "SKU", "Quantity", "Currency", "Amount FCY", "FX to AZN", "Base AZN", "Allocated Costs AZN", "Landed Value AZN", "Unit Cost AZN"],
    "Write-offs": ["Date", "Line ID", "Document", "SKU", "Product", "Quantity", "Unit Cost AZN", "Write-off Cost AZN"],
    "Checks": ["Severity", "Sheet", "Row", "ID", "Message"],
}


EMPTY_HEADERS['Shipment Cost'] += ['Supplier', 'Purchase Type']
EMPTY_HEADERS['Shipment Summary'] = ['Shipment ID', 'Purchase Type', 'Supplier', 'Base AZN', 'Allocated Costs AZN', 'Landed Value AZN']
EMPTY_HEADERS['Shipment Expenses'] = ['Shipment ID', 'Expense ID', 'Document Date', 'Expense Type', 'Counterparty', 'Document', 'Currency', 'Amount FCY', 'FX to AZN', 'Expense AZN']


def excel_value(value):
    return float(value) if isinstance(value, Decimal) else value


def write_sheet(wb, name, headers, records, note=""):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws["A2"] = name
    ws["A2"].font = Font(name="Arial", size=14, bold=True, color="222222")
    ws["A3"] = note
    ws["A3"].font = Font(name="Arial", size=10, color="63748A")
    ws.freeze_panes = "C7" if name == "Monthly Inventory" else "A7"
    for i, h in enumerate(headers, 1):
        cell = ws.cell(6, i, h)
        cell.font = Font(name="Arial", size=10, bold=True, color="222222")
        cell.fill = PatternFill("solid", fgColor="F2F2F2")
        cell.border = Border(bottom=Side(style='thin', color='999999'))
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = 22 if h in ("Product", "Customer", "Original Line ID") else 18
    ws.row_dimensions[6].height = 44
    body_font = Font(name="Arial", size=10, color="253448")
    body_alignment = Alignment(vertical="center")
    for ri, record in enumerate(records, 7):
        ws.row_dimensions[ri].height = 21
        for ci, h in enumerate(headers, 1):
            v = excel_value(record.get(h))
            cell = ws.cell(ri, ci, v)
            if isinstance(v, str):
                cell.data_type = "s"  # Never treat user text as a spreadsheet formula.
            cell.font = body_font
            cell.alignment = body_alignment
            if isinstance(v, (date, datetime)):
                cell.number_format = "mmm yyyy" if h == "Month" else "yyyy-mm-dd"
            elif isinstance(v, (int, float)):
                cell.number_format = "0.0%" if h == "Margin %" else "0.000000" if h in ("Unit Cost AZN", "Monthly AVCO AZN", "FX to AZN") else "#,##0.000;[Red](#,##0.000)" if ("Qty" in h or "Quantity" in h) else "0" if h == "Row" else "#,##0.00;[Red](#,##0.00)"
    last = max(7, 6 + len(records))
    # Excel forbids table names that are cell addresses (T1, T2, ...).
    # openpyxl writes such names without rejecting them; Excel repairs the file.
    table = Table(displayName="InventoryTable_" + str(len(wb.worksheets)), ref=f"A6:{get_column_letter(len(headers))}{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False, showColumnStripes=False)
    ws.add_table(table)
    if name == "Checks":
        ws.column_dimensions["E"].width = 95
        for row in ws.iter_rows(min_row=7, max_row=last):
            row[4].alignment = Alignment(wrap_text=True, vertical="center")
            ws.row_dimensions[row[0].row].height = 44
            if row[0].value == "ERROR":
                row[0].fill = PatternFill("solid", fgColor="FCE6E4")
                row[0].font = Font(name="Arial", size=10, bold=True, color="AA2929")
    return ws


def atomic_save(wb, output):
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=output.parent)
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, output)
    finally:
        wb.close()
        Path(tmp).unlink(missing_ok=True)


def build_report_workbook(result, input_path, digest, elapsed, language='en'):
    wb = Workbook()
    wb.remove(wb.active)
    cfg = result["config"]
    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
    for name, rows in result["tables"].items():
        note = f"{cfg['company']} | Through {cfg['end']:%Y-%m} | AZN excluding VAT | Generated {generated}"
        if name == 'Shipment Summary':
            note = 'Стоимость поступлений до конца отчётного месяца. Все известные расходы включены.'
        elif name == 'Shipment Expenses':
            note = 'Все расходы поставок из отчёта, включая поздние. Часть может относиться к поступлениям после отчётного месяца.'
        write_sheet(wb, name, EMPTY_HEADERS[name], rows, note)
    about = [
        {"Item": "Status", "Value": "SUCCESS — static results; run Recalculate again after changing input"},
        {"Item": "Company", "Value": cfg["company"]},
        {"Item": "Generated", "Value": generated}, {"Item": "Version", "Value": VERSION},
        {"Item": "Input file", "Value": Path(input_path).name}, {"Item": "Input SHA256", "Value": digest},
        {"Item": "Start Month", "Value": cfg["start"].isoformat()}, {"Item": "Report Through", "Value": cfg["end"].isoformat()},
        {"Item": "Calculation seconds", "Value": elapsed},
        {"Item": "Source counts", "Value": json.dumps(result["source_counts"], ensure_ascii=False)},
        {"Item": "Late expenses", "Value": "All loaded shipment expenses restate original receipt cost, regardless of expense document date."},
        {"Item": "Returns", "Value": "Customer returns: original sale month AVCO. Supplier returns: original purchase landed cost; this is inventory value, not supplier cash refund."},
        {"Item": "Profit", "Value": "Revenue and COGS are net of customer returns. Write-offs are separate. Gross profit less write-offs is not net profit."},
        {"Item": "Margin", "Value": "Gross profit / positive net revenue. Blank when net revenue is zero or negative."},
        {"Item": "Scope", "Value": "One pooled warehouse; zero opening stock at Start Month. Monthly negative balances checked; daily balances are not checked."},
        {"Item": "Precision", "Value": "Decimal calculations, no intermediate cent rounding. Excel displays money to 2 decimals and unit costs to 6."},
        {"Item": "Snapshots", "Value": "Sorting is safe. Filtering does not recalculate totals. Save a copy of this report if a historical snapshot is required."},
    ]
    ws = write_sheet(wb, "About", ["Item", "Value"], about, "How to read this report")
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 105
    for row in ws.iter_rows(min_row=7):
        row[1].alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[row[0].row].height = 38
    if language == 'az':
        from report_locale import localize_workbook
        localize_workbook(wb)
    return wb


def workbook_bytes(wb):
    """Serialize an openpyxl workbook without touching the host filesystem."""
    output = BytesIO()
    try:
        wb.save(output)
        return output.getvalue()
    finally:
        wb.close()


def write_report_bytes(result, input_path, digest, elapsed, language='en'):
    return workbook_bytes(build_report_workbook(result, input_path, digest, elapsed, language))


def write_report(result, output, input_path, digest, elapsed, language='en'):
    atomic_save(build_report_workbook(result, input_path, digest, elapsed, language), output)


def build_errors_workbook(issues, language='en'):
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "Checks", EMPTY_HEADERS["Checks"], issues, f"FAILED | {datetime.now(timezone.utc):%Y-%m-%d %H:%M UTC} | Inventory Report was not updated.")
    if language == 'az':
        from report_locale import localize_workbook
        localize_workbook(wb)
    return wb


def write_errors(issues, path, language='en'):
    atomic_save(build_errors_workbook(issues, language), path)


def write_errors_bytes(issues, language='en'):
    return workbook_bytes(build_errors_workbook(issues, language))


def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--prepare-input", action='store_true', help='Assign missing persistent row codes, backing up the input first.')
    parser.add_argument("--language", choices=('az', 'en'), default='az', help='Report language (default: az).')
    parser.add_argument("--input", type=Path, default=Path(__file__).with_name("Inventory_Input.xlsx"))
    parser.add_argument("--output", type=Path, default=Path(__file__).with_name("Inventory_Report.xlsx"))
    args = parser.parse_args(argv)
    src, dst = args.input.resolve(), args.output.resolve()
    if src == dst:
        print("Input and output must be different files.", file=sys.stderr)
        return 2
    log = dst.with_name(dst.stem + "_last_run.txt")
    checks = dst.with_name(dst.stem + "_Checks.xlsx")
    if src in (log, checks):
        print("Input must not use a generated log / checks filename.", file=sys.stderr)
        return 2
    started = time.perf_counter()
    try:
        raw, digest = read_input(src)
        if args.prepare_input:
            from prepare_input import prepare_input
            prepare_input(src)
            # Versioned books may refresh ranges without assigning any new codes.
            raw, digest = read_input(src)
        result = calculate(raw)
        calc_seconds = time.perf_counter() - started
        write_report(result, dst, src, digest, calc_seconds, args.language)
        message = f"SUCCESS | {datetime.now(timezone.utc).isoformat()}\nReport: {dst.name}\nInput SHA256: {digest}\nElapsed: {time.perf_counter()-started:.3f} seconds\n"
        warning_count = sum(x['Severity'] == 'WARNING' for x in result['tables']['Checks'])
        if warning_count:
            message += f"WARNING: {warning_count} possible duplicate(s). Review Checks.\n"
        try:
            log.write_text(message, encoding="utf-8")
        except OSError:
            message += "Report saved; the status log could not be written.\n"
        # An old error report is not deleted; its last-run timestamp identifies it.
        print(message)
        return 0
    except InvalidInput as e:
        message = f"FAILED | {datetime.now(timezone.utc).isoformat()}\n{len(e.issues)} error(s). Previous report NOT updated.\n"
        try:
            write_errors(e.issues, checks, args.language)
            message += f"See {checks.name}\n"
        except OSError:
            message += "Close the Checks workbook and retry.\n"
        message += "\n".join(f"{x['Sheet']} row {x['Row']}: {x['Message']}" for x in e.issues[:30])
    except PermissionError:
        message = "FAILED: Close the input and output workbooks in Excel and retry. Previous report NOT updated."
    except FileNotFoundError:
        message = f"FAILED: Input file not found: {src}"
    except Exception as e:
        message = f"FAILED: {type(e).__name__}: {e}. Previous report NOT updated."
    try:
        log.parent.mkdir(parents=True, exist_ok=True)
        log.write_text(message, encoding="utf-8")
    except OSError:
        pass
    print(message, file=sys.stderr)
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
