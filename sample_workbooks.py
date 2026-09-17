"""Public starter workbooks, generated only from code and fictional data."""
from datetime import date
from decimal import Decimal
from functools import lru_cache
from io import BytesIO
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from input_layout import AZ_SHEETS, FIELDS, INPUT_COLUMNS, SHEETS
from inventory import workbook_bytes
from migrate_input import migrate_input_bytes
from report_locale import LABELS, SHEETS as REPORT_SHEETS


TEMPLATE_FILE = 'downloads/Inventory_Template_AZ.xlsx'
DEMO_FILE = 'downloads/Inventory_Demo_AZ.xlsx'
CHECK_SHEET = 'Nümunə yoxlaması'

# Independently specified controls for the documented January–March example.
# table, month, SKU (None for monthly totals), metric, expected value
DEMO_CHECKS = (
    ('Monthly Inventory', 1, 'DEMO-A', 'Monthly AVCO AZN', '21'),
    ('Monthly Inventory', 1, 'DEMO-A', 'Closing Qty', '6'),
    ('Monthly Inventory', 1, 'DEMO-A', 'Closing Value AZN', '126'),
    ('Monthly Inventory', 2, 'DEMO-A', 'Monthly AVCO AZN', '21.625'),
    ('Monthly Inventory', 2, 'DEMO-A', 'Closing Qty', '11'),
    ('Monthly Inventory', 2, 'DEMO-A', 'Closing Value AZN', '237.875'),
    ('Monthly Inventory', 2, 'DEMO-B', 'Closing Qty', '2'),
    ('Monthly Inventory', 2, 'DEMO-B', 'Closing Value AZN', '42'),
    ('Monthly Summary', 2, None, 'Net Revenue AZN', '125'),
    ('Monthly Summary', 2, None, 'Net COGS AZN', '65.5'),
    ('Monthly Summary', 2, None, 'Gross Profit AZN', '59.5'),
    ('Monthly Summary', 2, None, 'Write-off Cost AZN', '21.625'),
    ('Monthly Inventory', 3, 'DEMO-A', 'Closing Value AZN', '237.875'),
    ('Monthly Inventory', 3, 'DEMO-B', 'Closing Value AZN', '42'),
)


def demo_rows():
    """Rows in INPUT_COLUMNS order. No local accounting workbook is read."""
    return {
        'Settings': [('Company', 'Nümunə şirkəti'),
                     ('Start Month', date(2025, 1, 1)),
                     ('Report Through', date(2025, 3, 1))],
        'Products': [('DEMO-A', 'Nümunə mal A', 'ədəd'),
                     ('DEMO-B', 'Nümunə mal B', 'ədəd')],
        'Counterparties': [(name,) for name in (
            'Nümunə təchizatçısı', 'Nümunə daşıyıcısı', 'Nümunə alıcı 1', 'Nümunə alıcı 2')],
        'Purchases': [
            (date(2025, 1, 10), 'Import', 'SHIP-01', 'Nümunə təchizatçısı', 'INV-001', 'DEMO-A', 10, 'USD', 100, 1.7, 'P1'),
            (date(2025, 1, 10), 'Import', 'SHIP-01', 'Nümunə təchizatçısı', 'INV-001', 'DEMO-B', 5, 'USD', 50, 1.7, 'P2'),
            (date(2025, 2, 5), 'Import', 'SHIP-02', 'Nümunə təchizatçısı', 'INV-002', 'DEMO-A', 10, 'EUR', 100, 2, 'P3'),
        ],
        'Expenses': [
            (date(2025, 3, 15), 'SHIP-01', 'Gömrük / broker', 'Nümunə daşıyıcısı', 'EXP-001', 'AZN', 60, 1, 'E1'),
            (date(2025, 2, 6), 'SHIP-02', 'Daşıma', 'Nümunə daşıyıcısı', 'EXP-002', 'AZN', 20, 1, 'E2'),
        ],
        'Sales': [
            (date(2025, 1, 20), 'SALE-001', 'Nümunə alıcı 1', 'DEMO-A', 'Sale', 4, 140, 'S1'),
            (date(2025, 1, 20), 'SALE-001', 'Nümunə alıcı 1', 'DEMO-B', 'Sale', 2, 60, 'S2'),
            (date(2025, 2, 15), 'SALE-002', 'Nümunə alıcı 2', 'DEMO-A', 'Sale', 5, 200, 'S3'),
            (date(2025, 2, 25), 'WO-001', None, 'DEMO-A', 'Write-off', 1, 0, 'W1'),
        ],
        'Returns': [
            (date(2025, 2, 16), 'Customer', 'S1', 1, 35, 'R1'),
            (date(2025, 2, 18), 'Customer', 'S3', 1, 40, 'R2'),
            (date(2025, 2, 20), 'Supplier', 'P1', 1, None, 'R3'),
            (date(2025, 2, 20), 'Supplier', 'P2', 1, None, 'R4'),
        ],
    }


def create_workbook(demo=False):
    rows = demo_rows() if demo else {
        'Settings': [(field, None) for field in ('Company', 'Start Month', 'Report Through')],
    }
    source = Workbook()
    source.remove(source.active)
    for name, columns in INPUT_COLUMNS.items():
        # Use the existing modern-input migration for the canonical AZ layout.
        ws = source.create_sheet(SHEETS[name])
        for ci, field in enumerate(columns, 1):
            ws.cell(6, ci, FIELDS[field])
        for ri, row in enumerate(rows.get(name, ()), 7):
            for ci, value in enumerate(row, 1):
                ws.cell(ri, ci, value)
    wb = load_workbook(BytesIO(migrate_input_bytes(workbook_bytes(source))))
    settings = wb[AZ_SHEETS['Settings']]
    for coordinate in ('B8', 'B9'):
        settings[coordinate].number_format = 'yyyy-mm-dd'
    if demo:
        add_demo_checks(wb)
    return workbook_bytes(wb)


def add_demo_checks(wb):
    ws = wb.create_sheet(CHECK_SHEET, 0)
    wb.active = 0
    ws.sheet_view.showGridLines = False
    ws['A2'] = 'Demo: məlumatlar və yoxlama'
    ws['A2'].font = Font(name='Arial', size=15, bold=True)
    notes = {
        3: 'Bütün şirkətlər və əməliyyatlar uydurmadır. Bu fayl yalnız sınaq üçündür.',
        4: 'Bu faylı saytda seçin, hesablayın və hesabatı endirin. Aşağıdakı nəticələri göstərilən vərəqlərlə müqayisə edin.',
        5: 'Nəticələr ilkin demo üçündür. Martda daxil edilən 60 AZN xərc yanvar alışlarına aiddir. Hesabat məbləğləri 2 onluq rəqəmə yuvarlaq göstərir.',
    }
    for ri, note in notes.items():
        ws.merge_cells(start_row=ri, start_column=1, end_row=ri, end_column=5)
        ws.cell(ri, 1, note).alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[ri].height = 36
    headers = ('Ay', 'Mal kodu (SKU)', 'Göstərici', 'Gözlənilən nəticə', 'Hesabat vərəqi')
    for ci, title in enumerate(headers, 1):
        cell = ws.cell(6, ci, title)
        cell.font = Font(name='Arial', bold=True)
        cell.fill = PatternFill('solid', fgColor='F2F2F2')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[6].height = 32
    for column, width in zip('ABCDE', (14, 19, 43, 25, 28)):
        ws.column_dimensions[column].width = width
    for ri, (table, month, sku, metric, expected) in enumerate(DEMO_CHECKS, 7):
        for ci, value in enumerate((date(2025, month, 1), sku or 'Bütün mallar',
                                    LABELS[metric], Decimal(expected), REPORT_SHEETS[table]), 1):
            ws.cell(ri, ci, value).alignment = Alignment(wrap_text=True, vertical='center')
        ws.cell(ri, 1).number_format = 'yyyy-mm'
        ws.cell(ri, 4).number_format = '#,##0.######'
        ws.row_dimensions[ri].height = 32
    ws.freeze_panes = 'A7'
    ws.auto_filter.ref = f'A6:E{ws.max_row}'


@lru_cache(maxsize=1)
def public_workbooks():
    return {TEMPLATE_FILE: create_workbook(), DEMO_FILE: create_workbook(demo=True)}


def workbook_parts(content):
    """Compare every XLSX member; only ZIP and document timestamps may vary."""
    with ZipFile(BytesIO(content)) as archive:
        names = archive.namelist()
        if len(names) != len(set(names)) or archive.comment:
            raise ValueError('Unexpected duplicate entries or archive comment')
        parts = {name: archive.read(name) for name in names}
    core = ET.fromstring(parts['docProps/core.xml'])
    for field in ('created', 'modified'):
        element = core.find(f'{{http://purl.org/dc/terms/}}{field}')
        if element is not None:
            element.text = ''
    parts['docProps/core.xml'] = ET.tostring(core)
    return parts
