"""Compatibility contracts: language/layout cannot change accounting results."""
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from openpyxl import Workbook, load_workbook
import inventory as inv
from input_layout import AZ_FIELDS, AZ_SHEETS, META_SHEET, schema_metadata
from migrate_input import migrate_input_bytes, refresh_ranges
from prepare_input import prepare_input_bytes
from report_locale import SHEETS as REPORT_SHEETS, LABELS
from test_inventory import fixture, write_modern_fixture


def serialize(wb):
    output = BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()


class MigrationTests(unittest.TestCase):
    def source(self, legacy=False):
        if legacy:
            wb = Workbook()
            wb.remove(wb.active)
            for name, rows in fixture().items():
                ws = wb.create_sheet(name)
                for ci, field in enumerate(inv.HEADERS[name], 1):
                    ws.cell(6, ci, field)
                for ri, row in enumerate(rows, 7):
                    for ci, field in enumerate(inv.HEADERS[name], 1):
                        ws.cell(ri, ci, row.get(field))
            return serialize(wb)
        with tempfile.TemporaryDirectory(dir=Path(__file__).parent / '.work') as tmp:
            path = Path(tmp) / 'input.xlsx'
            write_modern_fixture(path)
            return path.read_bytes()

    def test_ru_and_en_round_trip_preserve_all_calculations(self):
        for legacy in (False, True):
            with self.subTest(legacy=legacy):
                source = self.source(legacy)
                expected = inv.calculate(inv.read_input_bytes(source)[0])
                migrated = migrate_input_bytes(source)
                self.assertEqual(inv.calculate(inv.read_input_bytes(migrated)[0]), expected)
                again, assigned = prepare_input_bytes(migrated)
                self.assertEqual(assigned, 0)
                self.assertEqual(inv.calculate(inv.read_input_bytes(again)[0]), expected)
                self.assertEqual(inv.calculate(inv.read_input_bytes(migrate_input_bytes(again))[0]), expected)

    def test_codes_and_return_links_survive_sort_and_new_row(self):
        migrated = migrate_input_bytes(self.source())
        wb = load_workbook(BytesIO(migrated))
        ws = wb[AZ_SHEETS['Purchases']]
        rows = list(ws.iter_rows(min_row=7, max_row=9, values_only=True))
        for ri, row in enumerate(reversed(rows), 7):
            for ci, value in enumerate(row, 1):
                ws.cell(ri, ci, value)
        for ci, value in enumerate(rows[0], 1):
            ws.cell(12, ci, value)
        ws['K12'] = None
        prepared, count = prepare_input_bytes(serialize(wb))
        self.assertEqual(count, 1)
        parsed = inv.read_input_bytes(prepared)[0]
        self.assertEqual([p['Line ID'] for p in parsed['Purchases'][:3]], ['P3', 'P2', 'P1'])
        self.assertEqual(parsed['Returns'][0]['Original Line ID'], 'S1')
        self.assertEqual(parsed['Purchases'][3]['Line ID'], 'P000001')
        inv.calculate(parsed)

    def test_layout_lists_and_sparse_later_rows(self):
        wb = load_workbook(BytesIO(migrate_input_bytes(self.source())))
        try:
            self.assertEqual(schema_metadata(wb)['schema_version'], 2)
            self.assertEqual(wb[META_SHEET].sheet_state, 'hidden')
            self.assertTrue(all(not s.tables for s in wb))
            for name in ('Purchases', 'Sales', 'Returns', 'Expenses'):
                ws = wb[AZ_SHEETS[name]]
                self.assertEqual(ws.freeze_panes, 'A7')
                self.assertTrue(ws.auto_filter.ref.endswith('10006'))
                self.assertTrue(ws.data_validations.count)
                for rule in ws.data_validations.dataValidation:
                    self.assertNotIn('InputProducts[', rule.formula1)
            wb[AZ_SHEETS['Products']]['A11000'] = 'LATE-SKU'
            refresh_ranges(wb)
            self.assertIn('$A$12000', wb.defined_names['Inventory_SKUs'].attr_text)
            self.assertIn('12000', wb[AZ_SHEETS['Products']].auto_filter.ref)
            self.assertTrue(all(name in wb.defined_names for name in (
                'Inventory_SKUs', 'Inventory_Parties', 'Inventory_SalesCodes',
                'Inventory_PurchaseCodes', 'Inventory_Shipments')))
        finally:
            wb.close()

    def test_future_version_rejected_before_codes_assigned(self):
        wb = load_workbook(BytesIO(migrate_input_bytes(self.source())))
        wb[META_SHEET]['B1'] = 999
        data = serialize(wb)
        for operation in (inv.read_input_bytes, prepare_input_bytes, migrate_input_bytes):
            with self.assertRaises(inv.InvalidInput):
                operation(data)

    def test_conflicting_sheet_aliases_rejected(self):
        wb = load_workbook(BytesIO(self.source()))
        wb.create_sheet(AZ_SHEETS['Purchases'])
        data = serialize(wb)
        for operation in (inv.read_input_bytes, prepare_input_bytes, migrate_input_bytes):
            with self.assertRaises(inv.InvalidInput):
                operation(data)

    def test_migration_preserves_existing_data_errors(self):
        wb = load_workbook(BytesIO(self.source()))
        from input_layout import SHEETS
        ws = wb[SHEETS['Products']]
        ws['A10'], ws['B10'], ws['C10'] = 'DEMO-A', 'Duplicate', 'pcs'
        source = serialize(wb)
        messages = []
        for data in (source, migrate_input_bytes(source)):
            with self.assertRaises(inv.InvalidInput) as caught:
                inv.calculate(inv.read_input_bytes(data)[0])
            messages.append(caught.exception.issues)
        self.assertEqual(messages[0], messages[1])

    def test_az_report_values_equal_english_and_no_formulas_or_tables(self):
        result = inv.calculate(fixture())
        en = load_workbook(BytesIO(inv.write_report_bytes(result, 'input.xlsx', 'digest', 0)))
        az = load_workbook(BytesIO(inv.write_report_bytes(result, 'input.xlsx', 'digest', 0, 'az')))
        try:
            for name, records in result['tables'].items():
                ws = az[REPORT_SHEETS[name]]
                headers = inv.EMPTY_HEADERS[name]
                self.assertEqual([c.value for c in ws[6]], [LABELS[h] for h in headers])
                for ri in range(7, 7 + len(records)):
                    for ci, header in enumerate(headers, 1):
                        if header not in ('Type', 'Purchase Type', 'Severity', 'Sheet'):
                            self.assertEqual(ws.cell(ri, ci).value, en[name].cell(ri, ci).value)
            self.assertTrue(all(not ws.tables for ws in az))
            self.assertFalse(any(c.data_type == 'f' for ws in az for row in ws for c in row))
        finally:
            en.close()
            az.close()

    def test_duplicate_az_and_ru_headers_rejected(self):
        wb = load_workbook(BytesIO(self.source()))
        from input_layout import SHEETS
        wb[SHEETS['Products']]['D6'] = AZ_FIELDS['SKU']
        with self.assertRaises(inv.InvalidInput):
            inv.read_input_bytes(serialize(wb))

    def test_cli_digest_matches_refreshed_book_without_new_codes(self):
        from contextlib import redirect_stdout
        from io import StringIO
        from unittest.mock import patch
        import hashlib
        import prepare_input
        with tempfile.TemporaryDirectory(dir=Path(__file__).parent / '.work') as tmp:
            source = Path(tmp) / 'input.xlsx'
            output = Path(tmp) / 'report.xlsx'
            source.write_bytes(migrate_input_bytes(self.source()))
            original_prepare = prepare_input.prepare_input

            def refresh_book(path):
                self.assertEqual(original_prepare(path), 0)
                wb = load_workbook(path)
                wb.properties.description = 'Refreshed without new row codes'
                wb.save(path)
                wb.close()
                return 0

            with patch.object(prepare_input, 'prepare_input', side_effect=refresh_book), redirect_stdout(StringIO()):
                self.assertEqual(inv.main(['--input', str(source), '--output', str(output), '--prepare-input']), 0)
            wb = load_workbook(output)
            try:
                about = dict(wb[REPORT_SHEETS['About']].iter_rows(min_row=7, values_only=True))
                self.assertEqual(about[LABELS['Input SHA256']], hashlib.sha256(source.read_bytes()).hexdigest())
            finally:
                wb.close()


if __name__ == '__main__':
    unittest.main()
