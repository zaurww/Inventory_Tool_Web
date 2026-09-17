import copy
from datetime import date
from decimal import Decimal
import hashlib
from io import BytesIO
from pathlib import Path
import random
import re
import tempfile
import unittest
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from openpyxl import Workbook, load_workbook
import browser_api
import inventory as inv
from input_layout import SHEETS, FIELDS, INPUT_COLUMNS, SETTINGS, TYPES, canonical_columns
from input_layout import AZ_SHEETS
from prepare_input import prepare_input


def fixture():
    values = {
        "Settings": [["Company", "Demo Company"], ["Start Month", "2025-01-01"], ["Report Through", "2025-03-01"]],
        "Products": [["DEMO-A", "Product A", "pcs"], ["DEMO-B", "Product B", "pcs"]],
        "Purchases": [
            ["P1", "2025-01-10", "SHIP-01", "INV-001", "DEMO-A", 10, "USD", 100, 1.7],
            ["P2", "2025-01-10", "SHIP-01", "INV-001", "DEMO-B", 5, "USD", 50, 1.7],
            ["P3", "2025-02-05", "SHIP-02", "INV-002", "DEMO-A", 10, "EUR", 100, 2],
        ],
        "Expenses": [["E1", "2025-03-15", "SHIP-01", "Customs / broker", "AZN", 60, 1], ["E2", "2025-02-06", "SHIP-02", "Freight", "AZN", 20, 1]],
        "Sales": [
            ["S1", "2025-01-20", "SALE-001", "Customer One", "DEMO-A", "Sale", 4, 140],
            ["S2", "2025-01-20", "SALE-001", "Customer One", "DEMO-B", "Sale", 2, 60],
            ["S3", "2025-02-15", "SALE-002", "Customer Two", "DEMO-A", "Sale", 5, 200],
            ["W1", "2025-02-25", "WO-001", "", "DEMO-A", "Write-off", 1, 0],
        ],
        "Returns": [["R1", "2025-02-16", "Customer", "S1", 1, 35], ["R2", "2025-02-18", "Customer", "S3", 1, 40], ["R3", "2025-02-20", "Supplier", "P1", 1, None], ["R4", "2025-02-20", "Supplier", "P2", 1, None]],
    }
    return {name: [{**dict(zip(inv.HEADERS[name], row)), "_row": i + 7} for i, row in enumerate(rows)] for name, rows in values.items()}


def inventory_row(result, sku, m):
    return next(r for r in result["tables"]["Monthly Inventory"] if r["SKU"] == sku and r["Month"] == date(2025, m, 1))


class InventoryTests(unittest.TestCase):
    def test_exported_table_names_are_valid_excel_names(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / 'Report.xlsx'
            empty = fixture()
            for name in empty:
                if name != 'Settings':
                    empty[name] = []
            writers = [
                lambda: inv.write_report(inv.calculate(fixture()), output, 'Input.xlsx', 'test', 0),
                lambda: inv.write_report(inv.calculate(empty), output, 'Input.xlsx', 'test', 0),
                lambda: inv.write_errors([inv.issue('Purchases', 7, 'Test error')], output),
            ]
            for writer in writers:
                writer()
                with ZipFile(output) as archive:
                    names = set()
                    tables = [p for p in archive.namelist() if p.startswith('xl/tables/') and p.endswith('.xml')]
                    self.assertTrue(tables)
                    for part in tables:
                        table = ET.fromstring(archive.read(part))
                        name = table.attrib['displayName']
                        self.assertEqual(name, table.attrib['name'])
                        self.assertRegex(name, r'^[A-Za-z_][A-Za-z0-9_]*$')
                        self.assertFalse(re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]*', name, re.I), name)
                        self.assertFalse(re.fullmatch(r'R[1-9][0-9]*C[1-9][0-9]*', name, re.I), name)
                        self.assertNotIn(name.casefold(), {'r', 'c'} | names)
                        names.add(name.casefold())

    def test_golden_cost_and_profit(self):
        result = inv.calculate(fixture())
        january = inventory_row(result, "DEMO-A", 1)
        february = inventory_row(result, "DEMO-A", 2)
        march = inventory_row(result, "DEMO-A", 3)
        self.assertEqual(january["Monthly AVCO AZN"], Decimal("21"))
        self.assertEqual(january["Closing Value AZN"], Decimal("126"))
        self.assertEqual(february["Monthly AVCO AZN"], Decimal("21.625"))
        self.assertEqual(february["Closing Qty"], 11)
        self.assertEqual(february["Closing Value AZN"], Decimal("237.875"))
        self.assertEqual(february["Net Revenue AZN"], 125)
        self.assertEqual(february["Net COGS AZN"], Decimal("65.5"))
        self.assertEqual(february["Gross Profit AZN"], Decimal("59.5"))
        self.assertEqual(march["Opening Value AZN"], february["Closing Value AZN"])
        self.assertEqual(inventory_row(result, "DEMO-B", 2)["Closing Value AZN"], 42)
        for field in ("Net Revenue AZN", "Net COGS AZN", "Gross Profit AZN"):
            self.assertEqual(sum(r[field] for r in result["tables"]["Sales Detail"]), sum(r[field] for r in result["tables"]["Monthly Summary"]))

    def test_late_expense_restates_all_periods(self):
        raw = fixture()
        raw["Expenses"][0]["Amount FCY"] = 90
        result = inv.calculate(raw)
        self.assertEqual(inventory_row(result, "DEMO-A", 1)["Monthly AVCO AZN"], 23)
        self.assertEqual(inventory_row(result, "DEMO-A", 2)["Monthly AVCO AZN"], Decimal("22.375"))
        self.assertEqual(inventory_row(result, "DEMO-A", 3)["Closing Value AZN"], Decimal("246.125"))

    def test_same_month_return_does_not_change_avco(self):
        raw = fixture()
        raw["Returns"][1]["Quantity"] = 2
        raw["Returns"][1]["Revenue Reversal AZN"] = 80
        result = inv.calculate(raw)
        self.assertEqual(inventory_row(result, "DEMO-A", 2)["Monthly AVCO AZN"], Decimal("21.625"))
        self.assertEqual(inventory_row(result, "DEMO-A", 2)["Closing Qty"], 12)

    def test_sort_order_independent(self):
        raw = fixture()
        expected = inv.calculate(raw)
        for rows in raw.values():
            random.Random(42).shuffle(rows)
        self.assertEqual(inv.calculate(raw), expected)

    def test_repeated_purchase_sku_is_valid(self):
        self.assertEqual(len(inv.calculate(fixture())["tables"]["Shipment Cost"]), 3)

    def test_duplicate_line_id_is_error(self):
        raw = fixture()
        raw["Purchases"][2]["Line ID"] = "p1"
        with self.assertRaises(inv.InvalidInput):
            inv.calculate(raw)

    def test_duplicate_product_is_error(self):
        raw = fixture()
        raw["Products"].append({"SKU": "demo-a", "Product": "Duplicate", "Unit": "pcs", "_row": 9})
        with self.assertRaises(inv.InvalidInput):
            inv.calculate(raw)

    def test_missing_revenue_is_not_zero(self):
        raw = fixture()
        raw["Sales"][0]["Revenue AZN"] = None
        with self.assertRaises(inv.InvalidInput):
            inv.calculate(raw)

    def test_invalid_fx_and_unknown_shipment(self):
        for change in ("fx", "shipment"):
            raw = fixture()
            if change == "fx":
                raw["Purchases"][0]["FX to AZN"] = 0
            else:
                raw["Expenses"][0]["Shipment ID"] = "NOT-FOUND"
            with self.assertRaises(inv.InvalidInput):
                inv.calculate(raw)

    def test_over_return_and_return_before_sale(self):
        for field, value in [("Quantity", 5), ("Revenue Reversal AZN", 141), ("Date", "2025-01-01")]:
            raw = fixture()
            raw["Returns"][0][field] = value
            with self.assertRaises(inv.InvalidInput):
                inv.calculate(raw)

    def test_negative_inventory_is_error(self):
        raw = fixture()
        raw["Sales"][0]["Quantity"] = 100
        with self.assertRaises(inv.InvalidInput) as ctx:
            inv.calculate(raw)
        self.assertTrue(any(e["Sheet"] == "Monthly Inventory" for e in ctx.exception.issues))

    def test_cutoff_and_known_late_cost(self):
        raw = fixture()
        raw["Settings"][2]["Value"] = "2025-01-01"
        result = inv.calculate(raw)
        self.assertEqual(len(result["tables"]["Sales Detail"]), 2)
        self.assertEqual(inventory_row(result, "DEMO-A", 1)["Monthly AVCO AZN"], 21)
        self.assertEqual(len(result["tables"]["Shipment Cost"]), 2)

    def test_only_refund_month_has_blank_margin(self):
        raw = fixture()
        raw["Sales"] = [s for s in raw["Sales"] if s["Line ID"] != "S3"]
        raw["Returns"] = [r for r in raw["Returns"] if r["Return ID"] != "R2"]
        result = inv.calculate(raw)
        feb = next(r for r in result["tables"]["Monthly Summary"] if r["Month"] == date(2025, 2, 1))
        self.assertEqual(feb["Net Revenue AZN"], -35)
        self.assertIsNone(feb["Margin %"])

    def test_empty_input(self):
        raw = fixture()
        for name in raw:
            if name != "Settings":
                raw[name] = []
        self.assertEqual(inv.calculate(raw)["tables"]["Monthly Inventory"], [])

    def test_full_supplier_return_zero_balance(self):
        raw = fixture()
        raw["Sales"] = []
        raw["Returns"] = [{"Return ID": "R", "Date": "2025-01-20", "Type": "Supplier", "Original Line ID": "P1", "Quantity": 10, "Revenue Reversal AZN": None, "_row": 7}]
        result = inv.calculate(raw)
        self.assertEqual(inventory_row(result, "DEMO-A", 1)["Closing Qty"], 0)
        self.assertEqual(inventory_row(result, "DEMO-A", 1)["Closing Value AZN"], 0)

    def test_actual_workbook_and_no_formula_report(self):
        path = Path(__file__).parent / "examples" / "Inventory_Input_Demo.xlsx"
        if not path.exists():
            self.skipTest("Demo template not built yet")
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "Report.xlsx"
            self.assertEqual(inv.main(["--input", str(path), "--output", str(out)]), 0)
            wb = load_workbook(out, read_only=True, data_only=False)
            try:
                self.assertFalse(any(c.data_type == "f" for ws in wb for row in ws for c in row))
                self.assertEqual(wb["Aylıq anbar uçotu"]["H8"].value, 21.625)
            finally:
                wb.close()
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before)

    def test_failed_run_preserves_existing_report(self):
        path = Path(__file__).parent / "examples" / "Inventory_Input_Demo.xlsx"
        if not path.exists():
            self.skipTest("Demo template not built yet")
        with tempfile.TemporaryDirectory() as tmp:
            bad = Path(tmp) / "Bad.xlsx"
            wb = load_workbook(path)
            ws = wb[SHEETS['Sales']] if SHEETS['Sales'] in wb.sheetnames else wb['Sales']
            columns = canonical_columns([c.value for c in ws[6]])
            ws.cell(7, columns.index('Quantity') + 1, 100)
            wb.save(bad)
            wb.close()
            out = Path(tmp) / "Report.xlsx"
            out.write_bytes(b"previous-report-sentinel")
            self.assertEqual(inv.main(["--input", str(bad), "--output", str(out)]), 1)
            self.assertEqual(out.read_bytes(), b"previous-report-sentinel")
            self.assertTrue(out.with_name("Report_Checks.xlsx").exists())


def modern_fixture():
    raw = fixture()
    for p in raw['Purchases']:
        p.update(Supplier='Vendor', **{'Purchase Type': 'Import'})
    for e in raw['Expenses']:
        e.update(Counterparty='Carrier', Document=e['Expense ID'])
    raw['Counterparties'] = [{'Name': n, '_row': i + 7} for i, n in enumerate(('Vendor', 'Carrier', 'Customer One', 'Customer Two'))]
    return raw


def write_modern_fixture(path, blank_ids=False):
    raw = modern_fixture()
    wb = Workbook()
    wb.remove(wb.active)
    reverse_types = {v: k for k, v in TYPES.items()}
    for name, columns in INPUT_COLUMNS.items():
        ws = wb.create_sheet(SHEETS[name])
        for ci, field in enumerate(columns, 1):
            ws.cell(6, ci, FIELDS[field])
        for ri, row in enumerate(raw[name], 7):
            for ci, field in enumerate(columns, 1):
                value = row.get(field)
                if blank_ids and field in ('Line ID', 'Expense ID', 'Return ID'):
                    value = None
                if field == 'Setting':
                    value = SETTINGS.get(value, value)
                if field in ('Type', 'Purchase Type'):
                    value = reverse_types.get(value, value)
                ws.cell(ri, ci, value)
    wb.save(path)
    wb.close()


def write_browser_fixture(path):
    """Keep return source IDs stable while leaving unrelated new rows uncoded."""
    write_modern_fixture(path)
    wb = load_workbook(path)
    for sheet, cells in (
        (SHEETS['Purchases'], ('K9',)),
        (SHEETS['Expenses'], ('I7', 'I8')),
        (SHEETS['Sales'], ('H8', 'H10')),
        (SHEETS['Returns'], ('F7', 'F8', 'F9', 'F10')),
    ):
        for cell in cells:
            wb[sheet][cell] = None
    wb.save(path)
    wb.close()


class InputRedesignTests(unittest.TestCase):
    def test_browser_adapter_assigns_codes_and_returns_two_workbooks(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'Input.xlsx'
            write_browser_fixture(path)
            original = path.read_bytes()
            result = browser_api.process_workbook(original, path.name)
            self.assertTrue(result['ok'])
            self.assertEqual(result['assigned_codes'], 9)
            self.assertEqual(path.read_bytes(), original)
            updated = load_workbook(BytesIO(result['input_bytes']), read_only=True)
            report = load_workbook(BytesIO(result['report_bytes']), read_only=True, data_only=False)
            try:
                self.assertEqual(updated[AZ_SHEETS['Purchases']]['K9'].value, 'P000001')
                self.assertFalse(any(c.data_type == 'f' for ws in report for row in ws for c in row))
                self.assertEqual(report['Aylıq anbar uçotu']['H8'].value, 21.625)
            finally:
                updated.close()
                report.close()

    def test_browser_adapter_returns_checks_and_updated_codes_on_error(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'Input.xlsx'
            write_browser_fixture(path)
            wb = load_workbook(path)
            wb[SHEETS['Sales']]['F7'] = 100
            wb.save(path)
            wb.close()
            result = browser_api.process_workbook(path.read_bytes(), path.name)
            self.assertFalse(result['ok'])
            self.assertEqual(result['assigned_codes'], 9)
            self.assertTrue(result['issues'])
            updated = load_workbook(BytesIO(result['input_bytes']), read_only=True)
            checks = load_workbook(BytesIO(result['checks_bytes']), read_only=True)
            try:
                self.assertEqual(updated[AZ_SHEETS['Sales']]['H8'].value, 'S000001')
                self.assertEqual(checks['Yoxlamalar']['A7'].value, 'XƏTA')
            finally:
                updated.close()
                checks.close()

    def test_russian_workbook_preserves_monthly_results(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'Input.xlsx'
            write_modern_fixture(path)
            raw, _ = inv.read_input(path)
            result = inv.calculate(raw)
            expected = inv.calculate(fixture())
            for name in ('Monthly Inventory', 'Monthly Summary', 'Sales Detail', 'Write-offs'):
                self.assertEqual(result['tables'][name], expected['tables'][name])

    def test_duplicate_business_row_warns_but_is_counted(self):
        raw = fixture()
        p = copy.deepcopy(raw['Purchases'][0])
        p.update({'Line ID': 'P4', '_row': 10})
        raw['Purchases'].append(p)
        result = inv.calculate(raw)
        warnings = [r for r in result['tables']['Checks'] if r['Severity'] == 'WARNING']
        self.assertEqual(len(warnings), 1)
        self.assertEqual(warnings[0]['ID'], 'P4')
        self.assertEqual(inventory_row(result, 'DEMO-A', 1)['Purchases Qty'], 20)

    def test_duplicate_expense_warns_without_dropping_amount(self):
        raw = fixture()
        raw['Expenses'].append({**raw['Expenses'][0], 'Expense ID': 'E3', '_row': 9})
        result = inv.calculate(raw)
        self.assertTrue(any(r['Sheet'] == 'Expenses' and r['Severity'] == 'WARNING' for r in result['tables']['Checks']))
        summary = next(r for r in result['tables']['Shipment Summary'] if r['Shipment ID'] == 'SHIP-01')
        self.assertEqual(summary['Allocated Costs AZN'], 120)

    def test_counterparties_are_unique_and_references_are_checked(self):
        for mode in ('duplicate', 'missing'):
            raw = modern_fixture()
            if mode == 'duplicate':
                raw['Counterparties'].append({'Name': ' vendor ', '_row': 12})
            else:
                raw['Purchases'][0]['Supplier'] = 'Unknown'
            with self.assertRaises(inv.InvalidInput):
                inv.calculate(raw)

    def test_counterparty_spelling_is_canonicalized(self):
        raw = modern_fixture()
        raw['Purchases'][0]['Supplier'] = ' vendor '
        result = inv.calculate(raw)
        self.assertEqual(result['tables']['Shipment Cost'][0]['Supplier'], 'Vendor')

    def test_supplier_and_purchase_type_required_in_new_template(self):
        for field in ('Supplier', 'Purchase Type'):
            raw = modern_fixture()
            raw['Purchases'][0]['_modern'] = True
            raw['Purchases'][0][field] = ''
            with self.assertRaises(inv.InvalidInput):
                inv.calculate(raw)

    def test_inconsistent_shipment_type_is_error(self):
        raw = modern_fixture()
        raw['Purchases'][1]['Purchase Type'] = 'Local'
        with self.assertRaises(inv.InvalidInput):
            inv.calculate(raw)

    def test_shipment_reports_reconcile_and_include_late_expenses(self):
        raw = modern_fixture()
        raw['Settings'][2]['Value'] = '2025-01-01'
        result = inv.calculate(raw)
        summary = result['tables']['Shipment Summary']
        self.assertEqual(len(summary), 1)
        self.assertEqual(summary[0]['Base AZN'], 255)
        self.assertEqual(summary[0]['Allocated Costs AZN'], 60)
        self.assertEqual(summary[0]['Landed Value AZN'], 315)
        self.assertEqual(result['tables']['Shipment Expenses'][0]['Document Date'], date(2025, 3, 15))
        self.assertEqual(sum(r['Landed Value AZN'] for r in result['tables']['Shipment Cost']), summary[0]['Landed Value AZN'])

    def test_automatic_codes_back_up_and_remain_stable(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'Input.xlsx'
            write_modern_fixture(path, blank_ids=True)
            before = path.read_bytes()
            self.assertEqual(prepare_input(path), 13)
            backups = list((path.parent / 'backups').glob('*.xlsx'))
            self.assertEqual(len(backups), 1)
            self.assertEqual(backups[0].read_bytes(), before)
            after = path.read_bytes()
            self.assertEqual(prepare_input(path), 0)
            self.assertEqual(path.read_bytes(), after)
            wb = load_workbook(path)
            ws = wb[SHEETS['Purchases']]
            original = [[c.value for c in row] for row in ws.iter_rows(min_row=7, max_row=9)]
            for ri, row in enumerate(reversed(original), 7):
                for ci, value in enumerate(row, 1):
                    ws.cell(ri, ci, value)
            for ci, value in enumerate(original[0][:-1], 1):
                ws.cell(10, ci, value)
            wb.save(path)
            wb.close()
            self.assertEqual(prepare_input(path), 1)
            wb = load_workbook(path)
            ws = wb[SHEETS['Purchases']]
            self.assertEqual([ws.cell(i, 11).value for i in range(7, 10)], [r[-1] for r in reversed(original)])
            self.assertNotIn(ws.cell(10, 11).value, [r[-1] for r in original])
            wb.close()

    def test_existing_return_links_survive_preparation(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'Input.xlsx'
            write_modern_fixture(path)
            wb = load_workbook(path)
            wb[SHEETS['Expenses']]['I7'] = None
            wb.save(path)
            wb.close()
            self.assertEqual(prepare_input(path), 1)
            result = inv.calculate(inv.read_input(path)[0])
            self.assertEqual(inventory_row(result, 'DEMO-A', 2)['Net COGS AZN'], Decimal('65.5'))

    def test_missing_source_code_cannot_be_reused(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'Input.xlsx'
            write_modern_fixture(path, blank_ids=True)
            wb = load_workbook(path)
            wb[SHEETS['Returns']]['C7'] = 'S000001'
            wb.save(path)
            wb.close()
            prepare_input(path)
            wb = load_workbook(path)
            codes = [wb[SHEETS['Sales']].cell(i, 8).value for i in range(7, 11)]
            self.assertNotIn('S000001', codes)
            wb.close()

    def test_missing_counterparty_sheet_is_error_in_modern_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'Input.xlsx'
            write_modern_fixture(path)
            wb = load_workbook(path)
            del wb[SHEETS['Counterparties']]
            wb.save(path)
            wb.close()
            with self.assertRaises(inv.InvalidInput):
                inv.read_input(path)


if __name__ == "__main__":
    unittest.main()
