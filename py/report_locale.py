"""Presentation-only Azerbaijani report labels. Never translates business values."""
from math import ceil
from openpyxl.styles import Alignment
from input_layout import AZ_FIELDS, AZ_SETTINGS, AZ_SHEETS, AZ_TYPES

SHEETS = {
    **AZ_SHEETS, 'Monthly Summary': 'Aylıq yekun',
    'Monthly Inventory': 'Aylıq anbar uçotu', 'Sales Detail': 'Satışların təfərrüatı',
    'Sales Summary': 'Satışların yekunu', 'Shipment Cost': 'Tədarükün maya dəyəri',
    'Shipment Summary': 'Tədarüklərin yekunu', 'Shipment Expenses': 'Tədarük xərcləri',
    'Write-offs': 'Silinmələr', 'Checks': 'Yoxlamalar', 'About': 'Məlumat', 'All': 'Bütün məlumatlar',
}
LABELS = {
    **{k: v.rstrip(' *') for k, v in AZ_FIELDS.items()}, **AZ_SETTINGS,
    'Month': 'Ay', 'Opening Qty': 'İlkin qalıq miqdarı',
    'Opening Value AZN': 'İlkin qalıq dəyəri AZN', 'Purchases Qty': 'Alış miqdarı',
    'Purchases Value AZN': 'Alışların dəyəri AZN', 'Sales Qty': 'Satış miqdarı',
    'Monthly AVCO AZN': 'Aylıq orta maya dəyəri AZN',
    'Customer Return Qty': 'Alıcıdan qaytarma miqdarı',
    'Customer Return Value AZN': 'Alıcıdan qaytarma dəyəri AZN',
    'Supplier Return Qty': 'Təchizatçıya qaytarma miqdarı',
    'Supplier Return Value AZN': 'Təchizatçıya qaytarma dəyəri AZN',
    'Write-off Qty': 'Silinmə miqdarı', 'Write-off Cost AZN': 'Silinmənin maya dəyəri AZN',
    'Gross Sales COGS AZN': 'Satılan malların maya dəyəri AZN',
    'Net COGS AZN': 'Xalis satışların maya dəyəri AZN',
    'Closing Qty': 'Son qalıq miqdarı', 'Closing Value AZN': 'Son qalıq dəyəri AZN',
    'Net Revenue AZN': 'Xalis satış məbləği AZN', 'Gross Profit AZN': 'Ümumi mənfəət AZN',
    'Net Quantity': 'Xalis miqdar', 'Unit Cost AZN': 'Vahidin maya dəyəri AZN',
    'Margin %': 'Marja %', 'Gross Profit Less Write-offs AZN': 'Silinmələrdən sonra ümumi mənfəət AZN',
    'Base AZN': 'Alış dəyəri AZN', 'Allocated Costs AZN': 'Bölüşdürülmüş xərclər AZN',
    'Landed Value AZN': 'Tam maya dəyəri AZN', 'Expense AZN': 'Xərc AZN',
    'Severity': 'Səviyyə', 'Sheet': 'Vərəq', 'Row': 'Sətir', 'ID': 'Kod', 'Message': 'İzah',
    'Item': 'Göstərici', 'Status': 'Vəziyyət', 'Generated': 'Yaradılma vaxtı',
    'Version': 'Proqram versiyası', 'Input file': 'Giriş faylı',
    'Input SHA256': 'Giriş faylının SHA256 izi', 'Calculation seconds': 'Hesablama müddəti (saniyə)',
    'Source counts': 'Mənbə sətrlərinin sayı', 'Late expenses': 'Gec daxil edilmiş xərclər',
    'Returns': 'Qaytarmalar', 'Profit': 'Mənfəət', 'Margin': 'Marja',
    'Scope': 'Əhatə dairəsi', 'Precision': 'Dəqiqlik', 'Snapshots': 'Tarixi nüsxələr',
}
ABOUT_VALUES = {
    'Status': 'UĞURLU — statik nəticələr. Məlumat dəyişdikdə yenidən hesablayın.',
    'Late expenses': 'Sənəd tarixindən asılı olmayaraq bütün məlum xərclər ilkin alışın maya dəyərini yenidən hesablayır.',
    'Returns': 'Alıcı qaytarması ilkin satış ayının orta maya dəyəri ilə qiymətləndirilir. Təchizatçıya qaytarma ilkin alışın tam maya dəyəri ilə çıxılır; bu, pul geriödənişi deyil.',
    'Profit': 'Satış məbləği və maya dəyəri alıcı qaytarmaları çıxılmaqla göstərilir. Silinmələr ayrıdır. Silinmələrdən sonrakı ümumi mənfəət xalis mənfəət deyil.',
    'Margin': 'Ümumi mənfəət / müsbət xalis satış məbləği. Satış məbləği sıfır və ya mənfi olduqda boşdur.',
    'Scope': 'Bir ümumi anbar; başlanğıc qalıq sıfırdır. Mənfi qalıqlar gün üzrə deyil, ayın sonunda yoxlanılır.',
    'Precision': 'Hesablama Decimal ilə, aralıq yuvarlaqlaşdırma olmadan aparılır. Məbləğlər 2, vahid maya dəyəri 6 onluq rəqəmlə göstərilir.',
    'Snapshots': 'Filtrləmə yekunları yenidən hesablamır. Tarixi nəticələr üçün hesabatın ayrıca nüsxəsini saxlayın.',
}
SEVERITIES = {'ERROR': 'XƏTA', 'WARNING': 'XƏBƏRDARLIQ', 'INFO': 'MƏLUMAT', 'OK': 'UĞURLU'}


def localize_issues(issues):
    return [{**r, 'Sheet': SHEETS.get(r['Sheet'], r['Sheet'])} for r in issues]


def localize_workbook(wb):
    for ws in wb:
        name = ws.title
        headers = [c.value for c in ws[6]]
        ws.title = SHEETS[name]
        ws['A2'] = ws.title
        # Metadata strings only; never search/replace user-supplied cell contents.
        ws['A3'] = 'Statik hesabat. Məbləğlər AZN ilə, ƏDV-siz göstərilir.'
        if name == 'Checks':
            ws['A3'] = 'Xətaları və xəbərdarlıqları yoxlayın. Xəta olduqda yeni hesabat yaradılmır.'
        if name in ('Shipment Summary', 'Shipment Expenses'):
            ws['A3'] = 'Bütün məlum xərclər, o cümlədən gec xərclər daxil edilib. Xərclərin bir hissəsi gələcək alışlara aid ola bilər.'
        ws.row_dimensions[6].height = 64
        for ci, header in enumerate(headers, 1):
            ws.cell(6, ci, LABELS[header])
            widths = {'Product': 38, 'Supplier': 42, 'Customer': 36,
                      'Counterparty': 36, 'Shipment ID': 31, 'Document No': 25,
                      'SKU': 24, 'Expense Type': 28, 'Sheet': 28, 'Severity': 22}
            if header in widths:
                ws.column_dimensions[ws.cell(6, ci).column_letter].width = widths[header]
        for row in ws.iter_rows(min_row=7):
            if name == 'About':
                item = row[0].value
                if item in ABOUT_VALUES:
                    row[1].value = ABOUT_VALUES[item]
                elif item == 'Source counts':
                    import json
                    counts = json.loads(row[1].value)
                    row[1].value = json.dumps({SHEETS[k]: v for k, v in counts.items()}, ensure_ascii=False)
                row[0].value = LABELS.get(item, item)
            else:
                for ci, header in enumerate(headers):
                    cell = row[ci]
                    if header in ('Type', 'Purchase Type'):
                        cell.value = {**AZ_TYPES, 'Customer Return': 'Alıcıdan qaytarma'}.get(cell.value, cell.value)
                    elif name == 'Checks' and header == 'Sheet':
                        cell.value = SHEETS.get(cell.value, cell.value)
                    elif name == 'Checks' and header == 'Severity':
                        cell.value = SEVERITIES.get(cell.value, cell.value)
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                    cell.alignment = Alignment(wrap_text=True, vertical='center')
                    width = ws.column_dimensions[cell.column_letter].width or 18
                    lines = sum(max(1, ceil(len(line) / max(1, width - 3)))
                                for line in cell.value.split('\n'))
                    ws.row_dimensions[cell.row].height = max(
                        ws.row_dimensions[cell.row].height or 21, 16 * lines + 6)
                elif headers[cell.column - 1] == 'Month':
                    cell.number_format = 'yyyy-mm'
        for table in list(ws.tables):
            del ws.tables[table]
        ws.auto_filter.ref = f'A6:{ws.cell(6, len(headers)).column_letter}{max(7, ws.max_row)}'
    return wb
