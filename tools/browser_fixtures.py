"""Create synthetic browser QA workbooks, or verify downloaded results."""
from pathlib import Path
import argparse
import json
import sys
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import inventory as inv
from browser_api import process_workbook
from input_layout import SHEETS
from report_locale import SHEETS as REPORT_SHEETS
from test_inventory import write_modern_fixture


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('directory', type=Path)
    parser.add_argument('--verify', action='store_true')
    args = parser.parse_args()
    directory = args.directory.resolve()
    if not directory.is_relative_to(ROOT / '.work'):
        raise ValueError('QA artifacts must stay inside the project .work directory')
    if args.verify:
        expected = inv.read_input(directory / 'native-updated.xlsx')[0]
        downloaded = inv.read_input(directory / 'updated.xlsx')[0]
        repeated = inv.read_input(directory / 'repeat-updated.xlsx')[0]
        assert expected == downloaded == repeated
        native = load_workbook(directory / 'native-report.xlsx')
        browser = load_workbook(directory / 'report.xlsx')
        repeated_report = load_workbook(directory / 'repeat-report.xlsx')
        try:
            for name in native.sheetnames:
                if name == REPORT_SHEETS['About']:
                    continue  # Runtime, generated time and XLSX digest legitimately differ.
                rows = lambda wb: list(wb[name].iter_rows(min_row=6, values_only=True))
                assert rows(native) == rows(browser) == rows(repeated_report), name
            assert all(not ws.tables for ws in browser)
            assert not any(c.data_type == 'f' for ws in browser for row in ws for c in row)
        finally:
            native.close()
            browser.close()
            repeated_report.close()
        print('Browser and native Python: input data, stable codes and all report tables match.')
        return
    directory.mkdir(parents=True, exist_ok=False)
    source = directory / 'synthetic.xlsx'
    write_modern_fixture(source)
    wb = load_workbook(source)
    wb[SHEETS['Purchases']]['K9'] = None
    wb[SHEETS['Sales']]['H10'] = None
    wb.save(source)
    wb.close()
    result = process_workbook(source.read_bytes(), source.name)
    assert result['ok'] and result['assigned_codes'] == 2
    (directory / 'native-updated.xlsx').write_bytes(result['input_bytes'])
    (directory / 'native-report.xlsx').write_bytes(result['report_bytes'])
    wb = load_workbook(source)
    ws = wb[SHEETS['Products']]
    ws['A10'], ws['B10'], ws['C10'] = 'DEMO-A', 'Duplicate', 'pcs'
    wb.save(directory / 'invalid.xlsx')
    wb.close()
    print(json.dumps({'synthetic_fixture': str(source), 'assigned_codes': 2}))


if __name__ == '__main__':
    main()
