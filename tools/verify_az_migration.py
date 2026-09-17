"""Create a separate AZ copy and prove source/data/calculation preservation."""
from datetime import datetime
from io import BytesIO
import hashlib
import json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import browser_api
import inventory as inv
from input_layout import META_SHEET
from prepare_input import prepare_input_bytes
from openpyxl import load_workbook


def outcome(raw):
    try:
        return inv.calculate(raw)
    except inv.InvalidInput as exc:
        return exc.issues


def main():
    source = ROOT / 'Inventory_Input.xlsx'
    original = source.read_bytes()
    before_hash = hashlib.sha256(original).hexdigest()
    # Both paths assign the same new IDs; neither writes the user's file.
    prepared, assigned = prepare_input_bytes(original)
    baseline_raw = inv.read_input_bytes(prepared)[0]
    converted = browser_api.process_workbook(original, source.name)
    if not converted['migrated']:
        raise RuntimeError('Migration did not finish; no output saved.')
    converted_raw = inv.read_input_bytes(converted['input_bytes'])[0]
    assert baseline_raw == converted_raw, 'Canonical input differs'
    assert outcome(baseline_raw) == outcome(converted_raw), 'Calculation/validation differs'
    assert assigned == converted['assigned_codes']
    wb = load_workbook(BytesIO(converted['input_bytes']))
    try:
        assert wb[META_SHEET]['B1'].value == 2
        assert all(not ws.tables for ws in wb)
        assert not any(c.data_type == 'f' for ws in wb for row in ws for c in row)
    finally:
        wb.close()
    assert source.read_bytes() == original, 'Source changed during verification; retry with a saved book'
    output = ROOT / 'outputs' / ('az-v2-' + datetime.now().strftime('%Y%m%d-%H%M%S'))
    output.mkdir(parents=True, exist_ok=False)
    with (output / 'Inventory_Input_AZ.xlsx').open('xb') as target:
        target.write(converted['input_bytes'])
    filename = 'Inventory_Report_AZ.xlsx' if converted['ok'] else 'Inventory_Checks_AZ.xlsx'
    content = converted['report_bytes'] if converted['ok'] else converted['checks_bytes']
    with (output / filename).open('xb') as target:
        target.write(content)
    summary = dict(source_sha256=before_hash, source_unchanged=True,
                   canonical_data_equal=True, calculation_or_errors_equal=True,
                   assigned_codes=assigned, errors=len(converted['issues']),
                   output=str(output), report_ok=converted['ok'])
    print(json.dumps(summary, ensure_ascii=True, indent=2))


if __name__ == '__main__':
    main()
