"""Versioned AZ migration. Reads the source and only creates a NEW destination."""
from io import BytesIO
from pathlib import Path
import argparse
import math

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from input_layout import (
    AZ_FIELDS, AZ_SETTINGS, AZ_SHEETS, AZ_TYPES, INPUT_COLUMNS, META_SHEET,
    SCHEMA_VERSION, canonical_columns, canonical_setting, canonical_type,
    resolve_sheet, schema_metadata,
)

NOTES = {
    'Settings': 'Başlanğıc və son ayın ilk gününü göstərin. Başlanğıc qalıq sıfırdır.',
    'Products': 'Hər SKU yalnız bir dəfə yazılır. Malın adını və ölçü vahidini doldurun.',
    'Counterparties': 'Təchizatçıları, alıcıları və xidmət göstərənləri bir dəfə daxil edin.',
    'Purchases': 'Hər sətirdə bir mal. Məbləğ bütün sətrin ƏDV-siz alış məbləğidir.',
    'Expenses': 'Xərci bəyannamə / e-qaimə ilə bağlayın. Gec xərclər əvvəlki dövrləri yenidən hesablayır.',
    'Sales': 'Satış üçün alıcı və satış məbləği tələb olunur. Silinmədə məbləğ 0 və ya boşdur.',
    'Returns': 'İlkin əməliyyatın daimi kodunu seçin. Miqdar həmişə müsbətdir.',
}
WIDTHS = {'Date': 15, 'Document Date': 16, 'Shipment ID': 31, 'Supplier': 42,
          'Customer': 36, 'Counterparty': 36, 'Document': 25, 'SKU': 24,
          'Quantity': 17, 'Currency': 13, 'Amount FCY': 22, 'FX to AZN': 18,
          'Purchase Type': 18, 'Line ID': 19, 'Expense ID': 19, 'Return ID': 19,
          'Expense Type': 28, 'Type': 20, 'Revenue AZN': 23,
          'Original Line ID': 24, 'Revenue Reversal AZN': 28,
          'Product': 42, 'Unit': 20, 'Setting': 36, 'Value': 36, 'Name': 60}
GUIDE = [
    ('Məlumatların daxil edilməsi', 'Başlıqları dəyişməyin. Məlumatlar 7-ci sətirdən başlayır. Ulduz vacib sahəni göstərir.'),
    ('Yeni əməliyyat', 'Yeni əməliyyat üçün kodu boş saxlayın. Mövcud kodları dəyişməyin.'),
    ('Nəticə faylı', 'Hesablamadan sonra yenilənmiş giriş kitabını endirin və işi həmin kitabda davam etdirin.'),
    ('Qaytarma', 'Əvvəlcə ilkin əməliyyatı hesablayın, sonra onun koduna istinad edin.'),
    ('Siyahılar və filtrlər', 'Siyahıların və filtrlərin sərhədləri hər hesablamada yenilənir. Boş sətrlər siyahılarda görünə bilər.'),
    ('Sıralama', 'Bir sütunu ayrıca sıralamayın. Bütün məlumat diapazonunu birlikdə sıralayın.'),
    ('Yapışdırma', 'Formatı qorumaq üçün yalnız dəyərləri yapışdırın. Python yoxlamaları bütün dolu sətrləri oxuyur.'),
    ('Hesablama', 'Aylıq orta maya dəyəri (Monthly AVCO). Xərclər malların alış dəyərinə görə bölüşdürülür.'),
    ('Ehtiyat nüsxə', 'İlkin kitab avtomatik dəyişdirilmir. İş kitablarının və hesabatların ehtiyat nüsxələrini saxlayın.'),
]


def set_text(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = 's'


def refresh_ranges(wb):
    """Table-free lists work while editing Excel, with reserved space for new rows."""
    if not schema_metadata(wb):
        return
    sheets = {n: wb[resolve_sheet(wb, n)] for n in INPUT_COLUMNS}
    limits, columns = {}, {}
    for name, ws in sheets.items():
        columns[name] = canonical_columns([c.value for c in ws[6]])
        last = max((c.row for row in ws.iter_rows(min_row=7) for c in row
                    if c.value is not None), default=7)
        limits[name] = min(1048576, max(10006, last + 1000))
        ws.auto_filter.ref = f'A6:{get_column_letter(len(columns[name]))}{limits[name]}'
        ws.freeze_panes = 'A7'
    sources = {
        'Inventory_SKUs': ('Products', 'SKU'),
        'Inventory_Parties': ('Counterparties', 'Name'),
        'Inventory_Shipments': ('Purchases', 'Shipment ID'),
        'Inventory_PurchaseCodes': ('Purchases', 'Line ID'),
        'Inventory_SalesCodes': ('Sales', 'Line ID'),
    }
    for label, (name, field) in sources.items():
        col = get_column_letter(columns[name].index(field) + 1)
        ref = f'{quote_sheetname(sheets[name].title)}!${col}$7:${col}${limits[name]}'
        wb.defined_names.add(DefinedName(label, attr_text=ref))
    for name, ws in sheets.items():
        ws.data_validations.dataValidation.clear()
        for ci, field in enumerate(columns[name], 1):
            formula = None
            if field == 'SKU' and name != 'Products':
                formula = '=Inventory_SKUs'
            elif field in ('Supplier', 'Customer', 'Counterparty'):
                formula = '=Inventory_Parties'
            elif field == 'Shipment ID' and name == 'Expenses':
                formula = '=Inventory_Shipments'
            elif field == 'Purchase Type':
                formula = '"İdxal,Yerli"'
            elif field == 'Type':
                formula = '"Satış,Silinmə"' if name == 'Sales' else '"Alıcı,Təchizatçı"'
            elif field == 'Original Line ID':
                typ_col = get_column_letter(columns[name].index('Type') + 1)
                formula = f'IF(${typ_col}7="Alıcı",Inventory_SalesCodes,Inventory_PurchaseCodes)'
            if formula:
                rule = DataValidation(type='list', formula1=formula, allow_blank=True)
            elif field in ('Quantity', 'FX to AZN'):
                rule = DataValidation(type='decimal', operator='greaterThan', formula1=0, allow_blank=True)
            else:
                continue
            rule.showErrorMessage = True
            rule.errorStyle = 'stop'
            rule.errorTitle = 'Yanlış dəyər'
            rule.error = 'Siyahıdan dəyər seçin və ya tələb olunan müsbət ədədi daxil edin.'
            ws.add_data_validation(rule)
            col = get_column_letter(ci)
            rule.add(f'{col}7:{col}{limits[name]}')


def migrate_input_bytes(source):
    """Preserve source rows, values, identifiers and extra columns; never fix data."""
    from inventory import read_input_bytes, InvalidInput, issue, workbook_bytes
    raw, _ = read_input_bytes(source)  # Structural validation only; business errors remain visible.
    wb = load_workbook(BytesIO(source))
    try:
        meta = schema_metadata(wb)
        if meta:
            refresh_ranges(wb)
            return workbook_bytes(wb)
        # Renaming sheets cannot safely rewrite arbitrary custom formula dependencies.
        if any(c.data_type == 'f' for ws in wb for row in ws for c in row):
            raise InvalidInput([issue('All', 0, 'Avtomatik çevirmə üçün formulaları əvvəlcə dəyərlərlə əvəz edin. İlkin fayl dəyişməyib.')])
        if wb.defined_names or any(ws.defined_names for ws in wb):
            raise InvalidInput([issue('All', 0, 'Kitabda fərdi adlandırılmış diapazonlar var. Çevirmədən əvvəl onların əlaqələri yoxlanmalıdır.')])
        modern = any(r.get('_modern') for rows in raw.values() for r in rows)
        for name in INPUT_COLUMNS:
            title = resolve_sheet(wb, name)
            if title is None:  # Optional legacy counterparty directory.
                ws = wb.create_sheet(AZ_SHEETS[name])
                for ci, field in enumerate(INPUT_COLUMNS[name], 1):
                    set_text(ws.cell(6, ci), AZ_FIELDS[field])
            else:
                ws = wb[title]
                ws.title = AZ_SHEETS[name]
            columns = canonical_columns([c.value for c in ws[6]])
            for field in INPUT_COLUMNS[name]:
                if field not in columns:
                    columns.append(field)
            for ci, field in enumerate(columns, 1):
                if field in AZ_FIELDS:
                    set_text(ws.cell(6, ci), AZ_FIELDS[field])
            for row in ws.iter_rows(min_row=7):
                for ci, field in enumerate(columns):
                    cell = row[ci]
                    if cell.value is None:
                        continue
                    if field == 'Setting':
                        set_text(cell, AZ_SETTINGS.get(canonical_setting(cell.value), cell.value))
                    elif field in ('Type', 'Purchase Type'):
                        set_text(cell, AZ_TYPES.get(canonical_type(cell.value), cell.value))
            for table in list(ws.tables):
                del ws.tables[table]
            set_text(ws['A2'], AZ_SHEETS[name])
            set_text(ws['A3'], NOTES[name])
            set_text(ws['A4'], 'Sətir kodlarını saxlayın. Hesablamadan sonra yenilənmiş kitabda işləyin.')
            ws['A2'].font = Font(name='Arial', size=15, bold=True)
            for ri in (3, 4):
                ws.cell(ri, 1).alignment = Alignment(wrap_text=True, vertical='center')
                # Presentation note only; columns and data remain unmerged.
                if not any(ws.cell(ri, ci).value is not None for ci in range(2, len(columns)+1)):
                    for merged in list(ws.merged_cells.ranges):
                        if merged.min_row == ri and merged.max_row == ri:
                            ws.unmerge_cells(str(merged))
                    ws.merge_cells(start_row=ri, start_column=1, end_row=ri, end_column=len(columns))
                ws.row_dimensions[ri].height = 32
            ws.sheet_view.showGridLines = False
            ws.row_dimensions[6].height = 48
            for ci, field in enumerate(columns, 1):
                col = get_column_letter(ci)
                ws.column_dimensions[col].width = WIDTHS.get(field, 24)
                header = ws.cell(6, ci)
                header.font = Font(name='Arial', size=11, bold=True)
                header.fill = PatternFill('solid', fgColor='F2F2F2')
                header.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
                if field in ('SKU', 'Line ID', 'Expense ID', 'Return ID', 'Original Line ID', 'Document', 'Shipment ID'):
                    fmt = '@'
                elif field in ('Date', 'Document Date'):
                    fmt = 'yyyy-mm-dd'
                elif field in ('Quantity',):
                    fmt = '#,##0.000'
                elif field == 'FX to AZN':
                    fmt = '0.000000'
                elif field in ('Amount FCY', 'Revenue AZN', 'Revenue Reversal AZN'):
                    fmt = '#,##0.00'
                else:
                    fmt = None
                if fmt:
                    # Preserve values; reserve formatted cells for the next entries.
                    for ri in range(7, max(1006, ws.max_row) + 1):
                        ws.cell(ri, ci).number_format = fmt
            for ri in (3, 4):
                width = sum(WIDTHS.get(field, 24) for field in columns)
                lines = math.ceil(len(str(ws.cell(ri, 1).value)) / max(1, width * 0.85))
                ws.row_dimensions[ri].height = max(32, 15 * lines + 8)
            for row in ws.iter_rows(min_row=7):
                for ci, field in enumerate(columns):
                    cell = row[ci]
                    if cell.value is not None and field in ('Supplier', 'Customer', 'Counterparty', 'Product', 'Name', 'Document'):
                        cell.alignment = Alignment(wrap_text=True, vertical='center')
                        lines = math.ceil(len(str(cell.value)) / (WIDTHS.get(field, 24) * 0.9))
                        ws.row_dimensions[cell.row].height = max(
                            ws.row_dimensions[cell.row].height or 23, 15 * lines + 6)
        title = resolve_sheet(wb, 'Guide')
        if title:
            del wb[title]  # Replace program instructions, not accounting rows.
        guide = wb.create_sheet(AZ_SHEETS['Guide'])
        guide['A2'] = 'İstifadə qaydaları'
        guide['A2'].font = Font(name='Arial', size=15, bold=True)
        guide['A6'], guide['B6'] = 'Mövzu', 'Qayda'
        guide.column_dimensions['A'].width = 28
        guide.column_dimensions['B'].width = 100
        guide.sheet_view.showGridLines = False
        for ri, (label, note) in enumerate(GUIDE, 7):
            guide.cell(ri, 1, label)
            guide.cell(ri, 2, note).alignment = Alignment(wrap_text=True, vertical='center')
            guide.row_dimensions[ri].height = 42
        meta_ws = wb.create_sheet(META_SHEET)
        for row in [('schema_version', SCHEMA_VERSION), ('language', 'az'),
                    ('validation_profile', 'modern' if modern else 'legacy'),
                    ('validate_parties', 'Counterparties' in raw)]:
            meta_ws.append(row)
        meta_ws.sheet_state = 'hidden'
        refresh_ranges(wb)
        return workbook_bytes(wb)
    finally:
        wb.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', type=Path)
    parser.add_argument('destination', type=Path)
    args = parser.parse_args()
    if args.source.resolve() == args.destination.resolve() or args.destination.exists():
        parser.error('Destination must be a NEW file, different from the source.')
    output = migrate_input_bytes(args.source.read_bytes())
    args.destination.parent.mkdir(parents=True, exist_ok=True)
    with args.destination.open('xb') as target:
        target.write(output)
    print(args.destination)


if __name__ == '__main__':
    main()
