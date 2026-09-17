"""Public deployment boundary: only the explicit application allowlist."""
from pathlib import Path
from datetime import date
from decimal import Decimal
from io import BytesIO
import tempfile
import unittest
from openpyxl import load_workbook
from build_web import PUBLIC_FILES, validate_site
from inventory import VERSION, calculate, read_input_bytes, workbook_bytes
from input_layout import AZ_SHEETS, schema_metadata
from browser_api import process_workbook
from sample_workbooks import public_workbooks, TEMPLATE_FILE, DEMO_FILE, DEMO_CHECKS, CHECK_SHEET


class WebBoundaryTests(unittest.TestCase):
    def setUp(self):
        root = Path(__file__).parent / '.work'
        root.mkdir(exist_ok=True)
        self.tmp = tempfile.TemporaryDirectory(dir=root)
        self.addCleanup(self.tmp.cleanup)
        self.site = Path(self.tmp.name)
        for name in PUBLIC_FILES:
            target = self.site / name
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(public_workbooks().get(name, b''))
        (self.site / 'index.html').write_text(
            f'<span id="app-version">Inventory Tool {VERSION}</span>', encoding='utf-8')

    def test_complete_application_is_allowed(self):
        self.assertEqual(set(validate_site(self.site)), PUBLIC_FILES)

    def test_accounting_or_unexpected_file_blocks_publication(self):
        for name in ('Inventory_Input.xlsx', '.env', 'backups/source.xlsx'):
            with self.subTest(name=name):
                target = self.site / name
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_bytes(b'private data')
                with self.assertRaises(ValueError):
                    validate_site(self.site)
                target.unlink()

    def test_missing_module_blocks_publication(self):
        (self.site / 'py/inventory.py').unlink()
        with self.assertRaises(ValueError):
            validate_site(self.site)

    def test_replaced_public_workbook_blocks_publication(self):
        for name in (TEMPLATE_FILE, DEMO_FILE):
            with self.subTest(name=name):
                wb = load_workbook(BytesIO(public_workbooks()[name]))
                wb[AZ_SHEETS['Settings']]['B7'] = 'Private company data'
                (self.site / name).write_bytes(workbook_bytes(wb))
                with self.assertRaises(ValueError):
                    validate_site(self.site)
                (self.site / name).write_bytes(public_workbooks()[name])

    def test_outdated_ui_or_asset_version_blocks_publication(self):
        version = f'<span id="app-version">Inventory Tool {VERSION}</span>'
        for html in (version.replace(VERSION, '0.0.0'),
                     version + '<script src="./app.js?v=0.0.0"></script>'):
            with self.subTest(html=html):
                (self.site / 'index.html').write_text(html, encoding='utf-8')
                with self.assertRaises(ValueError):
                    validate_site(self.site)


class StarterWorkbookTests(unittest.TestCase):
    def test_template_has_blank_az_inputs_and_modern_validation(self):
        content = public_workbooks()[TEMPLATE_FILE]
        raw, _ = read_input_bytes(content)
        self.assertEqual([r['Setting'] for r in raw['Settings']],
                         ['Company', 'Start Month', 'Report Through'])
        self.assertTrue(all(r['Value'] is None for r in raw['Settings']))
        self.assertTrue(all(not rows for name, rows in raw.items() if name != 'Settings'))
        wb = load_workbook(BytesIO(content))
        try:
            self.assertEqual(schema_metadata(wb)['validation_profile'], 'modern')
            self.assertIs(schema_metadata(wb)['validate_parties'], True)
            self.assertTrue(all(not ws.tables for ws in wb))
            self.assertFalse(any(c.data_type == 'f' for ws in wb for row in ws for c in row))
            self.assertEqual(wb[AZ_SHEETS['Settings']]['B8'].number_format, 'yyyy-mm-dd')
            for name in ('Purchases', 'Expenses', 'Sales', 'Returns'):
                ws = wb[AZ_SHEETS[name]]
                self.assertTrue(ws.data_validations.count)
                self.assertEqual(ws.freeze_panes, 'A7')
                self.assertTrue(ws.auto_filter.ref.endswith('10006'))
        finally:
            wb.close()

    def test_filled_template_calculates_and_assigns_codes(self):
        wb = load_workbook(BytesIO(public_workbooks()[TEMPLATE_FILE]))
        for cell, value in (('B7', 'Synthetic company'), ('B8', date(2025, 1, 1)),
                            ('B9', date(2025, 1, 1))):
            wb[AZ_SHEETS['Settings']][cell] = value
        rows = {
            'Products': ('TEST-1', 'Test product', 'ədəd'),
            'Counterparties': ('Test vendor',),
            'Purchases': (date(2025, 1, 5), 'Yerli', 'TEST-INV', 'Test vendor',
                          'TEST-INV', 'TEST-1', 2, 'AZN', 30, 1),
        }
        for name, row in rows.items():
            for ci, value in enumerate(row, 1):
                wb[AZ_SHEETS[name]].cell(7, ci, value)
        result = process_workbook(workbook_bytes(wb), 'filled-template.xlsx')
        self.assertTrue(result['ok'], result.get('issues'))
        self.assertEqual(result['assigned_codes'], 1)
        report = calculate(read_input_bytes(result['input_bytes'])[0])
        stock = report['tables']['Monthly Inventory'][0]
        self.assertEqual(stock['Closing Qty'], 2)
        self.assertEqual(stock['Closing Value AZN'], 30)

    def test_demo_calculates_and_matches_published_controls(self):
        content = public_workbooks()[DEMO_FILE]
        processed = process_workbook(content, 'Inventory_Demo_AZ.xlsx')
        self.assertTrue(processed['ok'], processed.get('issues'))
        self.assertEqual(processed['warning_count'], 0)
        result = calculate(read_input_bytes(processed['input_bytes'])[0])
        wb = load_workbook(BytesIO(content))
        try:
            self.assertEqual(wb.active.title, CHECK_SHEET)
            for ri, (table, month, sku, metric, expected) in enumerate(DEMO_CHECKS, 7):
                row = next(r for r in result['tables'][table]
                           if r['Month'] == date(2025, month, 1)
                           and (sku is None or r['SKU'] == sku))
                self.assertEqual(row[metric], Decimal(expected), (table, month, sku, metric))
                self.assertEqual(Decimal(str(wb[CHECK_SHEET].cell(ri, 4).value)), Decimal(expected))
        finally:
            wb.close()


if __name__ == '__main__':
    unittest.main()
