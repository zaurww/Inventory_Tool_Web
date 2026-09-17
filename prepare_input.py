"""Assign persistent codes only to populated rows; back up before saving."""
from datetime import datetime
from io import BytesIO
from pathlib import Path
import shutil
from openpyxl import load_workbook
from input_layout import SHEETS, TYPES, canonical_columns
from input_layout import resolve_sheet, schema_metadata, canonical_type


def assign_persistent_codes(wb):
    from inventory import key, InvalidInput, issue
    meta = schema_metadata(wb)
    titles = {name: resolve_sheet(wb, name) for name in SHEETS if name != 'Guide'}
    changed = 0
    # A deleted source must never be silently replaced by a newly numbered row.
    reserved = {'Purchases': set(), 'Sales': set()}
    return_title = titles['Returns']
    if return_title:
        returns = wb[return_title]
        headers = canonical_columns([c.value for c in returns[6]])
        if 'Type' in headers and 'Original Line ID' in headers:
            for row in returns.iter_rows(min_row=7):
                typ = row[headers.index('Type')].value
                typ = canonical_type(typ)
                source = 'Sales' if key(typ) == 'customer' else 'Purchases' if key(typ) == 'supplier' else None
                if source:
                    reserved[source].add(key(row[headers.index('Original Line ID')].value))
    for name, field, prefix in (
        ('Purchases', 'Line ID', 'P'), ('Sales', 'Line ID', 'S'),
        ('Expenses', 'Expense ID', 'E'), ('Returns', 'Return ID', 'R'),
    ):
        title = titles[name]
        if title is None:
            continue
        ws = wb[title]
        columns = canonical_columns([c.value for c in ws[6]])
        if columns.count(field) != 1:
            raise InvalidInput([issue(name, 6, f'Sütun yoxdur və ya təkrarlanır: {field}.')])
        ci = columns.index(field) + 1
        used = {key(ws.cell(ri, ci).value) for ri in range(7, ws.max_row + 1)}
        used.update(reserved.get(name, set()))
        number = 1
        for row in ws.iter_rows(min_row=7):
            cell = row[ci - 1]
            if cell.value is not None and str(cell.value).strip():
                continue
            if not any(c.value is not None and str(c.value).strip() for c in row if c.column != ci):
                continue
            while key(f'{prefix}{number:06d}') in used:
                number += 1
            cell.value = f'{prefix}{number:06d}'
            cell.data_type = 's'
            used.add(key(cell.value))
            number += 1
            changed += 1
    if meta:
        from migrate_input import refresh_ranges
        refresh_ranges(wb)
    return changed


def prepare_input_bytes(raw_bytes):
    """Assign codes in memory and return (updated_bytes, changed_count)."""
    original = bytes(raw_bytes)
    wb = load_workbook(BytesIO(original))
    try:
        changed = assign_persistent_codes(wb)
        if not changed and not schema_metadata(wb):
            return original, 0
        output = BytesIO()
        wb.save(output)
        return output.getvalue(), changed
    finally:
        wb.close()


def prepare_input(path):
    from inventory import atomic_save
    path = Path(path)
    wb = load_workbook(path)
    try:
        changed = assign_persistent_codes(wb)
        if changed or schema_metadata(wb):
            # Never overwrite an earlier backup, even if two runs share a timestamp.
            backup_dir = path.parent / 'backups'
            backup_dir.mkdir(exist_ok=True)
            backup = backup_dir / f'{path.stem}_{datetime.now():%Y%m%d_%H%M%S_%f}.xlsx'
            with backup.open('xb') as target, path.open('rb') as source:
                shutil.copyfileobj(source, target)
            atomic_save(wb, path)
            print(f'Assigned {changed} row code(s). Backup: {backup}')
    finally:
        wb.close()
    return changed


if __name__ == '__main__':
    import argparse
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path)
    prepare_input(parser.parse_args().input)
